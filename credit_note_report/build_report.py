"""Builds Credit_Note_Report.docx from ebu_query_57ad.csv and the order-type breakdown."""
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from order_type_data import CREDIT_NOTES_COUNT, MONTHS, ORDER_TYPES

HERE = Path(__file__).parent
CHARTS = HERE / "charts"
CHARTS.mkdir(exist_ok=True)
OUT = HERE / "Credit_Note_Report.docx"

MONTH_FMT = "%b %Y"


def month_label(mm_yyyy: str) -> str:
    return pd.to_datetime("01-" + mm_yyyy, format="%d-%m-%Y").strftime(MONTH_FMT)


def pct(x: float, digits: int = 1) -> str:
    return f"{x:.{digits}f}%"


# ---------------------------------------------------------------- data
inv = pd.read_csv(HERE / "ebu_query_57ad.csv")
inv["RUN_DATE"] = pd.to_datetime(inv["RUN_DATE"], format="%d-%m-%Y")
inv = inv.sort_values("RUN_DATE", ascending=False).reset_index(drop=True)
inv["Month"] = inv["RUN_DATE"].dt.strftime(MONTH_FMT)
inv["CN_RATIO"] = inv["CREDIT_NOTES"] / inv["INVOICES"] * 100

ot = pd.DataFrame(ORDER_TYPES, index=[month_label(m) for m in MONTHS])
ot = ot.loc[:, ot.sum() > 0]
ot_share = ot.div(ot.sum(axis=1), axis=0) * 100
reported_total = pd.Series(CREDIT_NOTES_COUNT, index=ot.index)
gaps = (ot.sum(axis=1) - reported_total)
gaps = gaps[gaps != 0]

last12 = ot.iloc[:12]
year_share = (last12.sum() / last12.values.sum() * 100).sort_values(ascending=False)
year_window = f"{last12.index[-1]} – {last12.index[0]}"

# ---------------------------------------------------------------- charts
plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10})
NAVY, RED, GREY = "#1f3b73", "#c0392b", "#7f8c8d"


def chart_report1():
    d = inv
    fig, ax = plt.subplots(figsize=(10, 4.6))
    x = np.arange(len(d))
    w = 0.4
    ax.bar(x - w / 2, d["INVOICES"], w, color=NAVY, label="Invoices")
    ax.bar(x + w / 2, d["CREDIT_NOTES"], w, color=RED, label="Credit notes")
    for i, (a, b) in enumerate(zip(d["INVOICES"], d["CREDIT_NOTES"])):
        ax.text(i - w / 2, a, f"{a:,}", ha="center", va="bottom", fontsize=7, rotation=90)
        ax.text(i + w / 2, b, f"{b:,}", ha="center", va="bottom", fontsize=7, rotation=90)
    ax.set_xticks(x, d["Month"], rotation=45, ha="right")
    ax.set_ylabel("Count")
    ax.set_ylim(0, d["INVOICES"].max() * 1.3)
    ax2 = ax.twinx()
    ax2.plot(x, d["CN_RATIO"], color="#e67e22", marker="o", label="Credit notes per 100 invoices")
    ax2.set_ylabel("Credit notes as % of invoices")
    ax2.set_ylim(0, d["CN_RATIO"].max() * 1.4)
    h1, l1 = ax.get_legend_handles_labels()
    h2, l2 = ax2.get_legend_handles_labels()
    ax.legend(h1 + h2, l1 + l2, loc="upper left", fontsize=8, ncol=3)
    ax.set_title("Invoices vs credit notes by run month (latest on the left)")
    for s in ("top",):
        ax.spines[s].set_visible(False)
    fig.tight_layout()
    p = CHARTS / "report1_invoices_vs_credit_notes.png"
    fig.savefig(p, dpi=200)
    plt.close(fig)
    return p


def chart_report2():
    cols = ot_share.mean().sort_values(ascending=False).index
    data = ot_share[cols]
    fig, ax = plt.subplots(figsize=(11, 6.4))
    im = ax.imshow(data.values, cmap="YlOrRd", aspect="auto", vmin=0, vmax=60)
    ax.set_xticks(range(len(cols)), cols, rotation=55, ha="right", fontsize=8)
    ax.set_yticks(range(len(data)), data.index)
    for i in range(data.shape[0]):
        for j in range(data.shape[1]):
            v = data.values[i, j]
            if v == 0:
                label = "–"
            elif v < 0.1:
                label = "<0.1%"
            else:
                label = f"{v:.1f}%"
            ax.text(j, i, label, ha="center", va="center", fontsize=6.3,
                    color="white" if v > 35 else "black")
    ax.set_title("Share of credit notes by order type, per run month (latest at top)")
    cb = fig.colorbar(im, ax=ax, fraction=0.025, pad=0.01)
    cb.ax.yaxis.set_major_formatter(matplotlib.ticker.PercentFormatter(decimals=0))
    fig.tight_layout()
    p = CHARTS / "report2_order_type_share_by_month.png"
    fig.savefig(p, dpi=200)
    plt.close(fig)
    return p


def chart_report2_stacked():
    top = ot_share.mean().sort_values(ascending=False).index[:7]
    data = ot_share[top].copy()
    data["All other types"] = 100 - data.sum(axis=1)
    colors = ["#1f3b73", "#2e86c1", "#c0392b", "#e67e22", "#27ae60", "#8e44ad", "#16a085", "#bdc3c7"]
    fig, ax = plt.subplots(figsize=(10, 6))
    left = np.zeros(len(data))
    y = np.arange(len(data))
    for c, col in zip(colors, data.columns):
        vals = data[col].values
        ax.barh(y, vals, left=left, color=c, label=col, edgecolor="white", linewidth=0.5)
        for yi, (l, v) in enumerate(zip(left, vals)):
            if v >= 4:
                ax.text(l + v / 2, yi, f"{v:.0f}%", ha="center", va="center", fontsize=7, color="white")
        left += vals
    ax.set_yticks(y, data.index)
    ax.invert_yaxis()
    ax.set_xlim(0, 100)
    ax.xaxis.set_major_formatter(matplotlib.ticker.PercentFormatter())
    ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.06), ncol=4, fontsize=8, frameon=False)
    ax.set_title("Credit-note mix per run month: top 7 order types (latest at top)")
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    fig.tight_layout()
    p = CHARTS / "report2_order_type_mix_stacked.png"
    fig.savefig(p, dpi=200)
    plt.close(fig)
    return p


def chart_report3():
    d = year_share[year_share > 0].iloc[::-1]
    fig, ax = plt.subplots(figsize=(10, 6.4))
    colors = [RED if v >= 10 else NAVY if v >= 1 else GREY for v in d.values]
    ax.barh(d.index, d.values, color=colors)
    for i, v in enumerate(d.values):
        ax.text(v + 0.4, i, "<0.01%" if v < 0.01 else pct(v, 2 if v < 1 else 1), va="center", fontsize=8)
    ax.set_xlim(0, d.max() * 1.15)
    ax.xaxis.set_major_formatter(matplotlib.ticker.PercentFormatter())
    ax.set_title(f"Share of credit notes by order type, last 12 months ({year_window})")
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    fig.tight_layout()
    p = CHARTS / "report3_order_type_share_last_year.png"
    fig.savefig(p, dpi=200)
    plt.close(fig)
    return p


def chart_report3_donut():
    top = year_share.iloc[:4]
    data = pd.concat([top, pd.Series({"All other types": 100 - top.sum()})])
    colors = ["#1f3b73", "#c0392b", "#2e86c1", "#e67e22", "#bdc3c7"]
    fig, ax = plt.subplots(figsize=(7, 5))
    ax.pie(data.values, labels=[f"{k}\n{v:.1f}%" for k, v in data.items()], colors=colors,
           startangle=90, counterclock=False, wedgeprops=dict(width=0.38, edgecolor="white"),
           textprops=dict(fontsize=9))
    ax.set_title(f"Credit-note mix, last 12 months ({year_window})")
    fig.tight_layout()
    p = CHARTS / "report3_order_type_donut.png"
    fig.savefig(p, dpi=200)
    plt.close(fig)
    return p


# ---------------------------------------------------------------- doc helpers
doc = Document()
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(10.5)


def heading(text, level):
    h = doc.add_heading(text, level)
    for r in h.runs:
        r.font.color.rgb = RGBColor(0x1F, 0x3B, 0x73)
    return h


def para(text, bold_prefix=None, italic=False, size=None):
    p = doc.add_paragraph()
    if bold_prefix:
        p.add_run(bold_prefix).bold = True
    r = p.add_run(text)
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(4)
    return p


def bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        p.add_run(bold_prefix).bold = True
    p.add_run(text)
    p.paragraph_format.space_after = Pt(2)
    return p


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c, h in zip(t.rows[0].cells, headers):
        c.text = ""
        c.paragraphs[0].add_run(h).bold = True
    for row in rows:
        cells = t.add_row().cells
        for c, v in zip(cells, row):
            c.text = str(v)
    for row in t.rows:
        for i, c in enumerate(row.cells):
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
                if i > 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if widths:
                c.width = Inches(widths[i])
    return t


def picture(path, width=6.6):
    doc.add_picture(str(path), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


# ---------------------------------------------------------------- title
t = doc.add_heading("Invoices & Credit Notes Report", 0)
para(f"Run months covered: {inv['Month'].iloc[-1]} to {inv['Month'].iloc[0]} "
     f"({len(inv)} monthly runs). All months are ordered from the latest to the earliest.", italic=True)
para("Each run covers the billing cycle from the 2nd of the previous month up to the 1st of the "
     "run month (e.g. the Sep 2026 run covers 02 Aug 2026 – 01 Sep 2026).", italic=True)

heading("Executive summary", 1)
tot_inv, tot_cn = inv["INVOICES"].sum(), inv["CREDIT_NOTES"].sum()
hi = inv.loc[inv["CN_RATIO"].idxmax()]
lo = inv.loc[inv["CN_RATIO"].idxmin()]
bullet(f"{tot_inv:,} invoices and {tot_cn:,} credit notes were issued across the {len(inv)} runs, "
       f"i.e. {tot_cn / tot_inv * 100:.1f} credit notes for every 100 invoices.", "Overall: ")
bullet(f"{hi['Month']} had the highest credit-note rate ({hi['CN_RATIO']:.1f}% of invoices); "
       f"{lo['Month']} had the lowest ({lo['CN_RATIO']:.1f}%).", "Peak / low: ")
bullet(f"Over the last 12 months, “{year_share.index[0]}” orders account for "
       f"{pct(year_share.iloc[0])} of credit notes, followed by “{year_share.index[1]}” "
       f"({pct(year_share.iloc[1])}) and “{year_share.index[2]}” ({pct(year_share.iloc[2])}). "
       f"Together the top 3 are {pct(year_share.iloc[:3].sum())}.", "Order types: ")
big_months = inv[inv["INVOICES"] > 10000]["Month"].tolist()
bullet(f"Invoices spike to ~10.4k–10.8k every third month ({', '.join(big_months)}), "
       f"while the other months sit at ~5.4k–6.3k. Credit notes do not follow the same quarterly pattern.",
       "Pattern: ")

# ---------------------------------------------------------------- report 1
doc.add_page_break()
heading("Report 1 – Invoice count vs credit note count by month", 1)
picture(chart_report1())

rows = []
for i, r in inv.iterrows():
    prev = inv.iloc[i + 1] if i + 1 < len(inv) else None
    inv_chg = f"{(r['INVOICES'] / prev['INVOICES'] - 1) * 100:+.1f}%" if prev is not None else "–"
    cn_chg = f"{(r['CREDIT_NOTES'] / prev['CREDIT_NOTES'] - 1) * 100:+.1f}%" if prev is not None else "–"
    rows.append([r["Month"], f"{r['INVOICES']:,}", f"{r['CREDIT_NOTES']:,}", pct(r["CN_RATIO"]),
                 inv_chg, cn_chg])
rows.append(["Total", f"{tot_inv:,}", f"{tot_cn:,}", pct(tot_cn / tot_inv * 100), "", ""])
table(["Month", "Invoices", "Credit notes", "Credit notes / invoices", "Invoices MoM", "Credit notes MoM"],
      rows, [1.0, 0.9, 1.0, 1.4, 1.0, 1.2])
para("MoM = change versus the previous (earlier) month.", italic=True, size=8.5)

heading("Brief on each number", 2)
avg_inv, avg_cn = inv["INVOICES"].mean(), inv["CREDIT_NOTES"].mean()
for i, r in inv.iterrows():
    prev = inv.iloc[i + 1] if i + 1 < len(inv) else None
    s = (f"{r['INVOICES']:,} invoices and {r['CREDIT_NOTES']:,} credit notes, "
         f"{r['CN_RATIO']:.1f} credit notes per 100 invoices. ")
    if prev is not None:
        s += (f"Compared with {prev['Month']}, invoices moved {(r['INVOICES'] / prev['INVOICES'] - 1) * 100:+.1f}% "
              f"and credit notes {(r['CREDIT_NOTES'] / prev['CREDIT_NOTES'] - 1) * 100:+.1f}%. ")
    else:
        s += "This is the earliest run in the data, so it has no previous month to compare with. "
    if r["CREDIT_NOTES"] > avg_cn * 1.25:
        s += f"Credit notes are well above the {avg_cn:,.0f} monthly average."
    elif r["CREDIT_NOTES"] < avg_cn * 0.75:
        s += f"Credit notes are well below the {avg_cn:,.0f} monthly average."
    else:
        s += f"Credit notes are close to the {avg_cn:,.0f} monthly average."
    bullet(s, f"{r['Month']}: ")
bullet(f"{tot_inv:,} invoices (avg {avg_inv:,.0f}/month) and {tot_cn:,} credit notes "
       f"(avg {avg_cn:,.0f}/month) over {len(inv)} months, an overall rate of "
       f"{tot_cn / tot_inv * 100:.1f}%.", "Total: ")

# ---------------------------------------------------------------- report 2
doc.add_page_break()
heading("Report 2 – Percentage of credit notes per order type, by run month", 1)
para("Each row shows how that month's credit notes split across order types (every row adds up to 100%). "
     "Only percentages are shown. Order types that were zero in every month are left out.")
picture(chart_report2(), 6.9)
picture(chart_report2_stacked(), 6.6)

heading("Brief on each month", 2)
for m in ot_share.index:
    row = ot_share.loc[m].sort_values(ascending=False)
    row = row[row > 0]
    top3 = row.iloc[:3]
    s = "; ".join(f"{k} {pct(v)}" for k, v in top3.items())
    s += f". The top 3 make up {pct(top3.sum())}"
    rest = row.iloc[3:]
    small = rest[rest >= 1]
    if len(small):
        s += ". Other order types above 1%: " + ", ".join(f"{k} {pct(v)}" for k, v in small.items())
    tiny = rest[rest < 1]
    if len(tiny):
        s += f". {len(tiny)} more order type{'s' if len(tiny) > 1 else ''} each under 1% (together {pct(tiny.sum())})"
    s += "."
    bullet(s, f"{m}: ")

heading("Notable month-to-month shifts", 2)
for col in ["New", "Service Transfer", "Price Change", "Sale Suspension", "HW Replace"]:
    sr = ot_share[col]
    mx, mn = sr.idxmax(), sr.idxmin()
    bullet(f"ranges from {pct(sr.min())} ({mn}) to {pct(sr.max())} ({mx}); "
           f"its average monthly share is {pct(sr.mean())}.", f"{col}: ")

# ---------------------------------------------------------------- report 3
doc.add_page_break()
heading(f"Report 3 – Percentage of credit notes per order type, last year ({year_window})", 1)
para("“Last year” means the most recent 12 run months. All credit notes in that window are pooled and "
     "split by order type. Only percentages are shown.")
picture(chart_report3(), 6.6)
picture(chart_report3_donut(), 5.2)

heading("Brief on each order type", 2)
cum = 0
for k, v in year_share[year_share > 0].items():
    cum += v
    months_present = int((last12[k] > 0).sum())
    peak = ot_share.loc[last12.index, k].idxmax()
    share = "<0.01%" if v < 0.01 else pct(v, 2 if v < 1 else 1)
    s = (f"{share} of last year's credit notes (running total {pct(cum)}). "
         f"Appeared in {months_present} of 12 months; its highest monthly share was "
         f"{pct(ot_share.loc[peak, k])} in {peak}.")
    bullet(s, f"{k}: ")
zero_types = [k for k in ORDER_TYPES if k in last12 and last12[k].sum() == 0]
if zero_types:
    para(f"Order types with no credit notes last year: {', '.join(zero_types)}.", italic=True)

# ---------------------------------------------------------------- notes
heading("Data notes", 1)
bullet("Report 1 uses INVOICES and CREDIT_NOTES from ebu_query_57ad.csv.")
bullet("Reports 2 and 3 use the order-type breakdown from the screenshot (credit-note lines per order type). "
       "Percentages use each month's order-type total as 100%. This breakdown counts credit-note lines "
       "(CREDIT_NOTES_COUNT), which is a different measure from the CSV credit-note count used in Report 1.")
bullet("Always-zero order types are not charted: Add New Add on, Billing Suspension, Cancelation, "
       "Change Service Parameters, ESales/E-Sales Reactivation, Suspension and Termination, "
       "Hardware Change, Hardware Termination, Schedule Auto Termination, Service Change, Termination.")
if len(gaps):
    txt = ", ".join(f"{m} ({int(reported_total[m]):,} reported vs {int(ot.loc[m].sum()):,} summed)"
                    for m in gaps.index)
    bullet(f"The screenshot is low resolution. In these months the order types do not add up exactly to "
           f"the reported CREDIT_NOTES_COUNT: {txt}. The Sep 2026 and Mar 2026 totals are most likely "
           f"4,564 and 3,401 (a 4 read as a 6). If so, the Sep 2026 gap is 10 lines and the Mar 2026 gap "
           f"disappears. The remaining gaps of 2 to 20 lines move any share by less than 0.3 percentage points.")

doc.save(OUT)
print("wrote", OUT)
print(year_share.round(2).to_string())
if len(gaps):
    print("gaps:\n", gaps.to_string())
