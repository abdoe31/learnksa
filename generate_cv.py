#!/usr/bin/env python3
"""Generate Abdelrahman Hamdy CV (.docx) — YRP/Mrkbati removed; jobs & remaining projects expanded."""

from docx import Document
from docx.shared import Pt, RGBColor, Cm, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn, nsmap
from docx.oxml import OxmlElement

BLUE = RGBColor(0x2F, 0x54, 0x96)
GRAY = RGBColor(0x40, 0x40, 0x40)
BLACK = RGBColor(0x00, 0x00, 0x00)


def set_run(run, *, size=10, bold=False, color=None, name="Calibri"):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    if color is not None:
        run.font.color.rgb = color


def add_para(doc, text="", *, size=10, bold=False, color=None, space_before=0, space_after=4, align=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.08
    if text:
        run = p.add_run(text)
        set_run(run, size=size, bold=bold, color=color)
    return p


def add_labeled(doc, label, value, *, size=10, space_after=2):
    p = add_para(doc, size=size, space_after=space_after)
    r1 = p.add_run(label)
    set_run(r1, size=size, bold=True)
    r2 = p.add_run(value)
    set_run(r2, size=size, bold=False)
    return p


def add_bullet(doc, text, *, size=10):
    p = doc.add_paragraph(style="List Bullet")
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(2)
    pf.line_spacing = 1.08
    run = p.add_run(text)
    set_run(run, size=size, bold=False)
    return p


def add_heading_bar(doc, text):
    p = add_para(doc, text, size=11, bold=True, color=BLUE, space_before=10, space_after=4)
    # bottom border
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "2F5496")
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def add_job_header(doc, title, dates):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = True
    cell0, cell1 = table.rows[0].cells
    # clear default paragraphs
    for cell, text, align in (
        (cell0, title, WD_ALIGN_PARAGRAPH.LEFT),
        (cell1, dates, WD_ALIGN_PARAGRAPH.RIGHT),
    ):
        p = cell.paragraphs[0]
        p.alignment = align
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        set_run(run, size=10, bold=True)
    # remove table borders
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tblPr.append(borders)
    return table


def build():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(1.27)
        section.bottom_margin = Cm(1.27)
        section.left_margin = Cm(1.524)
        section.right_margin = Cm(1.524)
        section.page_width = Cm(21.59)
        section.page_height = Cm(27.94)

    # Header
    add_para(
        doc,
        "ABDELRAHMAN HAMDY ABDELRAHIM",
        size=16,
        bold=True,
        color=BLUE,
        space_after=2,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_para(
        doc,
        "Full Stack .NET | Angular Developer",
        size=11,
        bold=True,
        color=GRAY,
        space_after=2,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_para(
        doc,
        "Cairo, Egypt  |  +20 112 506 7726  |  abdelrahman.hamdy.abdelrahim@gmail.com",
        size=9,
        color=GRAY,
        space_after=0,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_para(
        doc,
        "GitHub: github.com/abdoe31  |  LinkedIn: linkedin.com/in/abdelrahman-hamdy-450a05208",
        size=9,
        color=GRAY,
        space_after=4,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )

    # Summary
    add_heading_bar(doc, "PROFESSIONAL SUMMARY")
    add_para(
        doc,
        "Full Stack .NET Developer with 3 years of experience building and maintaining production "
        "web applications using ASP.NET Core, Angular, SQL Server, and Entity Framework Core. "
        "Proven delivery of RESTful APIs, JWT authentication, Redis caching, and third-party "
        "integrations including payment gateways, government APIs, OpenAI agents, and WhatsApp Web. "
        "Focused on clean architecture, performance, and reliable enterprise/commercial systems.",
        size=10,
        space_after=4,
    )

    # Skills
    add_heading_bar(doc, "TECHNICAL SKILLS")
    add_labeled(doc, "Backend: ", "C#, ASP.NET Core (MVC & Web API), Entity Framework Core, SQL Server, LINQ, Redis")
    add_labeled(doc, "Frontend: ", "Angular, TypeScript, JavaScript, HTML5, CSS3, Bootstrap, Razor Pages")
    add_labeled(
        doc,
        "Integrations: ",
        "OpenAI Agents, WhatsApp Web, OPay, Arab Bank Payment Gateway, Government APIs, SMS Gateways",
    )
    add_labeled(
        doc,
        "Practices: ",
        "REST APIs, JWT Auth, Dependency Injection, Caching Strategies, SOLID, Git/GitHub",
        space_after=2,
    )

    # Experience
    add_heading_bar(doc, "PROFESSIONAL EXPERIENCE")

    add_job_header(doc, "Telecom Egypt — .NET Developer", "Feb 2025 – Present")
    add_para(doc, "Cairo, Egypt", size=9, color=GRAY, space_before=0, space_after=2)
    add_bullet(
        doc,
        "Develop and enhance enterprise billing systems using ASP.NET Core, SQL Server, and RESTful APIs for telecom customer and usage workflows.",
    )
    add_bullet(
        doc,
        "Design, maintain, and troubleshoot mission-critical backend services to support high-availability billing and operational processes.",
    )
    add_bullet(
        doc,
        "Implement and refine Entity Framework Core data access, stored-procedure/query paths, and API contracts used by internal enterprise consumers.",
    )
    add_bullet(
        doc,
        "Partner with cross-functional teams (business, QA, operations) to design, ship, and stabilize scalable features in production.",
    )

    add_job_header(doc, "TOQ for Information Technology — .NET Developer", "Aug 2024 – Jan 2025")
    add_para(doc, "Riyadh, Remote", size=9, color=GRAY, space_before=0, space_after=2)
    add_bullet(
        doc,
        "Delivered production ASP.NET Core and Angular applications for commercial clients, from API design through frontend consumption.",
    )
    add_bullet(
        doc,
        "Owned backend APIs, business workflows, and SQL Server / EF Core data models end to end — including validation, authorization, and release readiness.",
    )
    add_bullet(
        doc,
        "Integrated third-party services such as payment gateways, government APIs, OpenAI agents, and WhatsApp Web into client platforms.",
    )
    add_bullet(
        doc,
        "Improved application responsiveness with Redis caching strategies and SQL/query optimization on high-traffic endpoints.",
    )

    add_job_header(doc, "El Shorouk Academy — .NET Developer", "Aug 2023 – Oct 2024")
    add_para(doc, "Cairo, Egypt", size=9, color=GRAY, space_before=0, space_after=2)
    add_bullet(
        doc,
        "Built and maintained internal academic and administrative systems with ASP.NET Core, SQL Server, and Microsoft data technologies.",
    )
    add_bullet(
        doc,
        "Developed Web APIs and database-driven modules supporting enrollment, staff, and campus operational workflows.",
    )
    add_bullet(
        doc,
        "Implemented JWT authentication and role-based access patterns for secure multi-user campus applications.",
    )
    add_bullet(
        doc,
        "Maintained EF Core models, migrations, and SQL Server schemas; collaborated with stakeholders to iterate features in production use.",
    )

    # Projects (YRP & Mrkbati removed)
    add_heading_bar(doc, "KEY PROJECTS")

    add_para(doc, "Travego – Bus Booking System", size=10, bold=True, space_before=4, space_after=1)
    add_para(
        doc,
        "Online bus reservation platform covering trip search, seat booking, payments, and customer notifications.",
        size=9,
        color=GRAY,
        space_after=2,
    )
    add_bullet(
        doc,
        "Designed and implemented trip listing, seat selection, and reservation flows across ASP.NET Core Web API and Angular UI.",
    )
    add_bullet(
        doc,
        "Integrated OPay Payment Gateway for online checkout, payment confirmation, and booking status updates.",
    )
    add_bullet(
        doc,
        "Integrated SMS Gateway to send booking confirmations, reminders, and trip-related customer alerts.",
    )
    add_bullet(
        doc,
        "Applied Redis caching for trip schedules and availability lookups to reduce latency on high-read booking endpoints.",
    )
    add_bullet(
        doc,
        "Secured APIs with JWT authentication and structured SQL Server / EF Core models for trips, seats, and reservations.",
    )
    add_labeled(
        doc,
        "Tech: ",
        "ASP.NET Core, Angular, SQL Server, EF Core, Redis, OPay, SMS, JWT",
        size=9,
        space_after=4,
    )

    add_para(doc, "TheMatrix – E-learning Platform", size=10, bold=True, space_before=2, space_after=1)
    add_para(
        doc,
        "Learning Management System for educational institutions — courses, assessments, grading, and multi-role access.",
        size=9,
        color=GRAY,
        space_after=2,
    )
    add_bullet(
        doc,
        "Built course management features including video lectures, quizzes, assignments, submissions, and grading workflows.",
    )
    add_bullet(
        doc,
        "Implemented JWT-secured, role-based access for students, instructors, and admins across backend APIs and Angular screens.",
    )
    add_bullet(
        doc,
        "Delivered RESTful ASP.NET Core Web APIs consumed by an Angular frontend for catalog, enrollment, and assessment flows.",
    )
    add_bullet(
        doc,
        "Modeled courses, enrollments, and assessments with Entity Framework Core and SQL Server for reliable academic data management.",
    )
    add_labeled(
        doc,
        "Tech: ",
        "ASP.NET Core Web API, Angular, SQL Server, EF Core, JWT",
        size=9,
        space_after=4,
    )

    # Education
    add_heading_bar(doc, "EDUCATION")
    add_job_header(doc, "Bachelor of Computer Science — El Shorouk Academy", "2018 – 2022")
    add_para(doc, "Graduation Project: Excellent", size=9, color=GRAY, space_after=2)
    add_job_header(doc, "Full Stack Web Development (.NET & Angular) — ITI", "2023")
    add_para(
        doc,
        "Information Technology Institute — intensive full stack program",
        size=9,
        color=GRAY,
        space_after=2,
    )

    out = "Abdelrahman_Hamdy_CV.docx"
    doc.save(out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    build()
