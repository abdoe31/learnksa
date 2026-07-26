#!/usr/bin/env python3
"""Generate a concise Arabic RTL Word CV for Mahmoud Hamed."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ACCENT = RGBColor(0x0B, 0x4F, 0x6C)
INK = RGBColor(0x1A, 0x23, 0x32)
MUTED = RGBColor(0x5C, 0x6B, 0x7A)


def set_run_font(run, size=11, bold=False, color=INK, name="Arial"):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:cs"), name)
    rfonts.set(qn("w:eastAsia"), name)


def set_paragraph_rtl(paragraph, space_after=6, space_before=0, align=WD_ALIGN_PARAGRAPH.RIGHT):
    paragraph.alignment = align
    pf = paragraph.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    ppr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    ppr.append(bidi)


def add_horizontal_line(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "18")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "0B4F6C")
    p_bdr.append(bottom)
    ppr.append(p_bdr)


def add_heading_ar(doc, text):
    p = doc.add_paragraph()
    set_paragraph_rtl(p, space_after=8, space_before=12)
    run = p.add_run(text)
    set_run_font(run, size=13, bold=True, color=ACCENT)
    add_horizontal_line(p)
    return p


def add_text(doc, text, size=11, bold=False, color=INK, space_after=6, space_before=0, align=WD_ALIGN_PARAGRAPH.RIGHT):
    p = doc.add_paragraph()
    set_paragraph_rtl(p, space_after=space_after, space_before=space_before, align=align)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    set_paragraph_rtl(p, space_after=3)
    if p.runs:
        p.runs[0].text = text
        set_run_font(p.runs[0], size=11, color=INK)
    else:
        run = p.add_run(text)
        set_run_font(run, size=11, color=INK)
    return p


def set_document_rtl(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:cs"), "Arial")
    sect_pr = doc.sections[0]._sectPr
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    sect_pr.append(bidi)


def build_cv(output_path: str):
    doc = Document()
    set_document_rtl(doc)

    section = doc.sections[0]
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    add_text(
        doc,
        "محمود محمد عبدالفتاح حامد",
        size=20,
        bold=True,
        color=ACCENT,
        space_after=4,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_text(
        doc,
        "Dispatcher",
        size=12,
        bold=True,
        color=INK,
        space_after=4,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    contact = add_text(
        doc,
        "01102600149  |  hhuumahmoud@gmail.com  |  القاهرة",
        size=10,
        color=MUTED,
        space_after=4,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_horizontal_line(contact)

    add_heading_ar(doc, "نبذة")
    add_text(
        doc,
        "شغال Dispatcher في شركات توصيل. متعود أوزع أوردرات، أتابع المناديب، وأحل مشاكل الشيفت بسرعة.",
        size=11,
        space_after=6,
    )

    add_heading_ar(doc, "الخبرة")

    add_text(doc, "Dispatcher — Breadfast", size=12, bold=True, space_after=1)
    add_text(doc, "يناير 2026 – يونيو 2026", size=10, color=MUTED, space_after=3)
    for item in [
        "توزيع الأوردرات على المناديب حسب المنطقة والوقت.",
        "متابعة حالة التوصيل والتدخل لو حصل تأخير.",
        "التنسيق مع الدعم والتشغيل وقت الضغط.",
    ]:
        add_bullet(doc, item)

    add_text(doc, "Dispatcher — InstaShop", size=12, bold=True, space_before=8, space_after=1)
    add_text(doc, "2024 – 2025", size=10, color=MUTED, space_after=3)
    for item in [
        "تنظيم شغل المناديب ومتابعة الأوردرات اليومية.",
        "التعامل مع مشاكل التوصيل أثناء الشيفت.",
        "التواصل مع العملاء عند الحاجة بخصوص حالة الطلب.",
    ]:
        add_bullet(doc, item)

    add_heading_ar(doc, "التعليم")
    add_text(doc, "بكالوريوس علوم حاسب — المعهد العالي بالتجمع الخامس", size=12, bold=True, space_after=1)
    add_text(doc, "متوقع التخرج 2027", size=10, color=MUTED, space_after=6)

    add_heading_ar(doc, "المهارات")
    for item in [
        "توزيع ومتابعة الأوردرات",
        "العمل تحت الضغط",
        "التواصل مع المناديب والعملاء",
        "Microsoft Office / Excel",
        "العربية، وإنجليزي متوسط",
    ]:
        add_bullet(doc, item)

    doc.save(output_path)
    print(f"Saved: {output_path}")


if __name__ == "__main__":
    build_cv("/workspace/CV-محمود-محمد-عبدالفتاح-حامد.docx")
