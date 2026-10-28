"""Builds the PLC study presentation (Netcracker Product Management, idea to CRM).

Run:  pip install python-pptx  &&  python3 build_deck.py
"""

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).parent
IMG = HERE / "images"
OUT = HERE / "PLC_Netcracker_Presentation.pptx"

PURPLE = RGBColor(0x5B, 0x2C, 0x83)
LILAC = RGBColor(0xEE, 0xE7, 0xF5)
DARK = RGBColor(0x26, 0x22, 0x2E)
GREY = RGBColor(0x5F, 0x5A, 0x6B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0xE0, 0x7A, 0x1F)
GREEN = RGBColor(0x2E, 0x8B, 0x57)
RED = RGBColor(0xB0, 0x3A, 0x2E)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
W, H = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]
TOTAL = 17
_page = [0]


def box(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    s.shadow.inherit = False
    return s


def text(slide, x, y, w, h, content, size=16, color=DARK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, shape=None):
    """content: str, or list of str / (bold_lead, rest) tuples, one paragraph each."""
    s = shape or slide.shapes.add_textbox(x, y, w, h)
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    items = content if isinstance(content, list) else [content]
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if isinstance(it, tuple):
            r = p.add_run()
            r.text = it[0]
            r.font.bold = True
            r.font.size = Pt(size)
            r.font.color.rgb = PURPLE if color == DARK else color
            r2 = p.add_run()
            r2.text = it[1]
            r2.font.size = Pt(size)
            r2.font.color.rgb = color
        else:
            r = p.add_run()
            r.text = it
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
        if len(items) > 1:
            p.space_after = Pt(size * 0.55)
    return s


def bullets(slide, x, y, w, h, items, size=17):
    s = slide.shapes.add_textbox(x, y, w, h)
    tf = s.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(9)
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(Inches(0.3)))
        pPr.set("indent", str(-Inches(0.25)))
        bu = pPr.makeelement(qn("a:buChar"), {"char": "•"})
        pPr.append(bu)
        lead, rest = it if isinstance(it, tuple) else ("", it)
        if lead:
            r = p.add_run()
            r.text = lead
            r.font.bold = True
            r.font.size = Pt(size)
            r.font.color.rgb = PURPLE
        r = p.add_run()
        r.text = rest
        r.font.size = Pt(size)
        r.font.color.rgb = DARK
    return s


def slide(title, subtitle=None):
    _page[0] += 1
    s = prs.slides.add_slide(BLANK)
    box(s, 0, 0, W, Inches(0.12), fill=PURPLE)
    text(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.8), title, size=30, color=PURPLE, bold=True)
    if subtitle:
        text(s, Inches(0.6), Inches(1.05), Inches(12), Inches(0.5), subtitle, size=16, color=GREY)
    text(s, Inches(0.6), Inches(7.0), Inches(9), Inches(0.35),
         "PLC study  |  Netcracker Product Management: from idea to CRM", size=10, color=GREY)
    text(s, Inches(11.7), Inches(7.0), Inches(1.0), Inches(0.35), f"{_page[0]} / {TOTAL}", size=10,
         color=GREY, align=PP_ALIGN.RIGHT)
    return s


def picture(s, file, x, y, max_w, max_h, caption=None):
    iw, ih = Image.open(IMG / file).size
    scale = min(max_w / iw, max_h / ih)
    w, h = int(iw * scale), int(ih * scale)
    px = x + (max_w - w) // 2
    box(s, px - Inches(0.04), y - Inches(0.04), w + Inches(0.08), h + Inches(0.08), line=RGBColor(0xD0, 0xC6, 0xDD))
    s.shapes.add_picture(str(IMG / file), px, y, w, h)
    if caption:
        text(s, x, y + h + Inches(0.08), max_w, Inches(0.5), caption, size=10, color=GREY,
             align=PP_ALIGN.CENTER)


def arrow(s, x1, y1, x2, y2, color=PURPLE):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.line.color.rgb = color
    c.line.width = Pt(2)
    ln = c.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle"}))
    return c


def table(s, x, y, w, rows, col_w, size=13, row_h=Inches(0.42)):
    shp = s.shapes.add_table(len(rows), len(rows[0]), x, y, w, row_h * len(rows))
    t = shp.table
    for c, cw in enumerate(col_w):
        t.columns[c].width = Inches(cw)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            cell.text = val
            cell.margin_left = cell.margin_right = Inches(0.08)
            p = cell.text_frame.paragraphs[0]
            p.runs[0].font.size = Pt(size)
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = PURPLE
                p.runs[0].font.bold = True
                p.runs[0].font.color.rgb = WHITE
            else:
                cell.fill.fore_color.rgb = LILAC if r % 2 else WHITE
                p.runs[0].font.color.rgb = DARK
                p.runs[0].font.bold = c == 0
    return t


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t


# ------------------------------------------------------------------ 1 title
_page[0] += 1
s = prs.slides.add_slide(BLANK)
box(s, 0, 0, W, H, fill=PURPLE)
box(s, Inches(0.8), Inches(2.2), Inches(0.12), Inches(2.4), fill=ACCENT)
text(s, Inches(1.1), Inches(2.1), Inches(11), Inches(0.6), "PLC STUDY", size=16, color=LILAC, bold=True)
text(s, Inches(1.1), Inches(2.6), Inches(11.5), Inches(1.6),
     "Product Lifecycle Management in Netcracker", size=40, color=WHITE, bold=True)
text(s, Inches(1.1), Inches(3.6), Inches(11), Inches(0.8), "From idea to CRM: how a telecom-grade PLC works",
     size=24, color=LILAC)
text(s, Inches(1.1), Inches(5.6), Inches(11), Inches(0.5), "Prepared by: [Your name]   |   October 2026",
     size=14, color=LILAC)
notes(s, "Purpose: walk through how a proven telecom PLC system takes a product from idea to the CRM, "
         "using Netcracker as the reference system.")

# ------------------------------------------------------------------ 2 agenda
s = slide("Agenda")
items = [
    ("01", "Why a PLC system", "What it gives an operator"),
    ("02", "PLC and TM Forum in brief", "The industry lifecycle model"),
    ("03", "Netcracker at a glance", "Modules and data model"),
    ("04", "Using the system", "Eight stages from idea to launch-ready"),
    ("05", "Integration with CRM", "How the product reaches the salesperson"),
    ("06", "Evidence", "Results at real operators"),
]
for i, (n, t, d) in enumerate(items):
    col, row = i % 2, i // 2
    x, y = Inches(0.8 + col * 6.1), Inches(1.6 + row * 1.7)
    box(s, x, y, Inches(5.7), Inches(1.4), fill=LILAC)
    text(s, x + Inches(0.2), y + Inches(0.2), Inches(1.1), Inches(1.0), n, size=36, color=PURPLE, bold=True)
    text(s, x + Inches(1.3), y + Inches(0.25), Inches(4.3), Inches(0.5), t, size=20, color=DARK, bold=True)
    text(s, x + Inches(1.3), y + Inches(0.75), Inches(4.3), Inches(0.5), d, size=14, color=GREY)

# ------------------------------------------------------------------ 3 takeaways
s = slide("Why a PLC system: what it gives an operator")
cards = [
    ("One source of truth", "Every product, price and rule is defined once. CRM, ordering and billing read it "
     "instead of keeping their own copies."),
    ("A governed journey", "Fixed statuses and approval gates. Nothing reaches sales until it is approved "
     "and tested."),
    ("Business users configure", "New offers are built from reusable templates and components, without "
     "development."),
    ("Rules travel with the product", "Eligibility by market, channel and customer category is enforced "
     "automatically when sales quotes."),
    ("Standard integration", "The catalog is exposed through TM Forum Open APIs, so any CRM can consume it."),
    ("Proven results", "Zain KSA: 50% decrease in product development time."),
]
for i, (t, d) in enumerate(cards):
    col, row = i % 3, i // 3
    x, y = Inches(0.6 + col * 4.1), Inches(1.5 + row * 2.65)
    box(s, x, y, Inches(3.8), Inches(2.4), fill=LILAC)
    box(s, x, y, Inches(3.8), Inches(0.1), fill=ACCENT if i == 5 else PURPLE)
    text(s, x + Inches(0.15), y + Inches(0.3), Inches(3.5), Inches(0.6), t, size=19, color=PURPLE, bold=True)
    text(s, x + Inches(0.15), y + Inches(0.9), Inches(3.5), Inches(1.4), d, size=15, color=DARK)
notes(s, "These six points are what the rest of the presentation shows in practice.")

# ------------------------------------------------------------------ 4 TMF620 lifecycle
s = slide("PLC and the TM Forum standard in brief",
          "TMF620 (Product Catalog Management) defines the statuses every catalog item goes through")
stages = [("In Study", "Idea studied"), ("In Design", "Being designed"), ("In Test", "Being tested"),
          ("Active", "Ready, not yet sold"), ("Launched", "On sale"), ("Retired", "No new sales")]
gates = ["Conception\napproved", "Design\napproved", "Test\napproved", "Start of\nmarketing", "End of\nmarketing"]
bw, gap = Inches(1.75), Inches(0.42)
x0, y0 = Inches(0.6), Inches(2.6)
for i, (st, d) in enumerate(stages):
    x = x0 + i * (bw + gap)
    fill = PURPLE if i < 4 else (ACCENT if i == 4 else GREY)
    b = box(s, x, y0, bw, Inches(1.1), fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, 0, 0, 0, 0, [st], size=18, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, shape=b)
    text(s, x, y0 + Inches(1.2), bw, Inches(0.5), d, size=13, color=GREY, align=PP_ALIGN.CENTER)
    if i < 5:
        arrow(s, x + bw + Inches(0.03), y0 + Inches(0.55), x + bw + gap - Inches(0.03), y0 + Inches(0.55))
        text(s, x + bw - Inches(0.45), y0 - Inches(0.85), gap + Inches(0.9), Inches(0.7), gates[i], size=11,
             color=PURPLE, align=PP_ALIGN.CENTER)
rx = x0 + 2 * (bw + gap)
b = box(s, rx, Inches(4.55), bw, Inches(0.6), fill=RED, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 0, 0, 0, 0, ["Rejected (test KO)"], size=13, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
     anchor=MSO_ANCHOR.MIDDLE, shape=b)
box(s, x0 - Inches(0.1), Inches(5.55), 4 * (bw + gap) - gap + Inches(0.2), Inches(0.08), fill=PURPLE)
text(s, x0, Inches(5.7), Inches(8), Inches(0.5), "Before CRM: defined, approved and tested inside the catalog",
     size=14, color=PURPLE, bold=True)
text(s, x0 + 4 * (bw + gap), Inches(5.7), Inches(4), Inches(0.5), "CRM: product is visible to sales",
     size=14, color=ACCENT, bold=True)
text(s, x0, Inches(6.3), Inches(12), Inches(0.5),
     "Netcracker is TM Forum-certified on TMF620 (June 2024), so its lifecycle follows this industry model.",
     size=13, color=GREY)

# ------------------------------------------------------------------ 5 modules
s = slide("Netcracker Product Management at a glance",
          "Two core components: Product Catalog + Product Lifecycle Management")
picture(s, "01_tmforum_netcracker_module_map.png", Inches(0.6), Inches(1.7), Inches(5.6), Inches(4.4),
        "Source: TM Forum SID certification report, Netcracker module map")
table(s, Inches(6.6), Inches(1.7), Inches(6.2), [
    ["Module", "What it does"],
    ["Product Lifecycle Mgmt", "Concept to launch; versions; approvals"],
    ["Product Offering Catalog", "Offers, bundles, rules, prices"],
    ["Technical Product Catalog", "Specifications, delivery templates"],
    ["Templates & components", "Fast, reusable configuration"],
    ["Pricing", "Recurring, one-time, discounts, taxes"],
    ["Business rules", "Eligibility, dependencies, constraints"],
    ["Catalog federation", "Partner catalogs in one place"],
    ["Open API layer", "Exposes the catalog to CRM and others"],
], [2.6, 3.6], size=13)

# ------------------------------------------------------------------ 6 catalog split
s = slide("Commercial and technical catalogs",
          "What is sold is kept separate from how it is delivered, and both feed sales and orders")
picture(s, "02_netcracker_catalog_feeds_sales_and_orders.png", Inches(0.6), Inches(1.7), Inches(7.4),
        Inches(4.8), "Source: Netcracker CPQ whitepaper (2021), Figure 2")
for i, (t, d) in enumerate([
    ("Product Offering Catalog", "The commercial view: offering details, business rules, prices and discounts."),
    ("Technical Product Catalog", "The technical view: product and service specifications, provisioning "
     "templates and decomposition rules."),
    ("Why it matters", "Marketing can create many offers on the same technical product without changing "
     "the network side."),
]):
    y = Inches(1.8 + i * 1.6)
    box(s, Inches(8.4), y, Inches(4.4), Inches(1.4), fill=LILAC)
    text(s, Inches(8.55), y + Inches(0.1), Inches(4.1), Inches(0.4), t, size=16, color=PURPLE, bold=True)
    text(s, Inches(8.55), y + Inches(0.5), Inches(4.1), Inches(0.9), d, size=13, color=DARK)

# ------------------------------------------------------------------ 7 data model
s = slide("Core data model, with one running example",
          "TM Forum-certified entities in Netcracker, illustrated with \"Business Internet 100\"")
model = [("Product Specification", "What it technically is", "Dedicated Internet Access, 100 Mbps"),
         ("Product Offering", "What is sold (single or bundle)", "\"Business Internet 100\" bundle"),
         ("Product Offering Price", "Prices attached to the offer", "Monthly fee + installation per site"),
         ("Product Catalog", "Grouping of offerings", "Enterprise connectivity catalog"),
         ("Product", "What the customer owns after sale", "Company X, 3 sites")]
for i, (e, m, ex) in enumerate(model):
    y = Inches(1.7 + i * 0.95)
    b = box(s, Inches(0.8), y, Inches(3.6), Inches(0.8), fill=PURPLE if i < 4 else GREY,
            shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, 0, 0, 0, 0, [e], size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, shape=b)
    text(s, Inches(4.6), y + Inches(0.18), Inches(4.0), Inches(0.5), m, size=15, color=DARK)
    b2 = box(s, Inches(8.6), y, Inches(4.1), Inches(0.8), fill=LILAC, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, 0, 0, 0, 0, [ex], size=14, color=PURPLE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=b2)
text(s, Inches(8.6), Inches(1.35), Inches(4.1), Inches(0.4), "Example (illustrative)", size=12, color=GREY,
     align=PP_ALIGN.CENTER)
text(s, Inches(0.8), Inches(6.55), Inches(12), Inches(0.4),
     "\"Business Internet 100\" = 100 Mbps dedicated internet + static IPs + managed router + SLA, "
     "on 12 / 24 / 36-month contracts.", size=13, color=GREY)


# ------------------------------------------------------------------ 8-9 stages
def stage_slide(title, subtitle, cards):
    s = slide(title, subtitle)
    for i, (num, name, status, owner, steps) in enumerate(cards):
        x = Inches(0.6 + i * 3.1)
        y = Inches(1.65)
        box(s, x, y, Inches(2.9), Inches(5.15), fill=LILAC)
        hb = box(s, x, y, Inches(2.9), Inches(1.0), fill=PURPLE)
        text(s, 0, 0, 0, 0, [f"Stage {num}", name], size=16, color=WHITE, bold=True, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.MIDDLE, shape=hb)
        text(s, x + Inches(0.1), y + Inches(1.1), Inches(2.7), Inches(0.35), status, size=12, color=ACCENT,
             bold=True)
        text(s, x + Inches(0.1), y + Inches(1.42), Inches(2.7), Inches(0.5), "Owner: " + owner, size=12,
             color=GREY)
        bullets(s, x + Inches(0.05), y + Inches(2.0), Inches(2.8), Inches(3.1), steps, size=13)
    return s


s = stage_slide("Using the system (1/2): define the product",
                "What the user does in the system, stage by stage (screen names are illustrative)", [
    ("1", "Ideation & business case", "In Study", "Product Manager", [
        "Register the concept \"Business Internet 100 for SMEs\"",
        "Attach business case: segment, sites, revenue",
        "Use consumption analytics to size demand",
        "Reuse existing catalog items",
        "Submit for conception approval"]),
    ("2", "Technical definition", "In Design", "Product Specialist", [
        "Pick the \"Dedicated Internet Access\" spec",
        "Set bandwidth, access type, IP block",
        "Add router, IP block, SLA monitoring specs",
        "Define decomposition: survey, install, configure"]),
    ("3", "Commercial offer", "In Design", "Product Manager", [
        "Create the offering from a template",
        "Bundle internet + IPs + router + SLA",
        "Add optional add-ons (extra IPs, backup link)",
        "Place it in the Enterprise catalog"]),
    ("4", "Pricing", "In Design", "PM + Finance", [
        "Monthly fee and installation per site",
        "Prices by contract term (12 / 24 / 36)",
        "Multi-site discount, taxes",
        "Price validity period"]),
])
notes(s, "Evidence: Netcracker PLM covers marketing concept, configuration, testing and launch; catalog "
         "covers specifications, offerings, bundles and pricing.")

s = stage_slide("Using the system (2/2): control and release it",
                "Rules, approvals, versions and testing decide when a product is ready", [
    ("5", "Rules & eligibility", "In Design", "Product Manager", [
        "Market: fiber-covered areas",
        "Channel: enterprise account managers",
        "Customer category: SME, corporate",
        "Dependencies and constraints (backup link only on 24/36 months)"]),
    ("6", "Approval gates", "In Design to In Test", "Approvers", [
        "PM submits for design approval",
        "Marketing, Finance, Legal / Regulatory, IT review",
        "Approve or reject with comments",
        "All approved: status moves to In Test"]),
    ("7", "Versioning", "Any status", "PM / Catalog admin", [
        "Never edit an approved product directly",
        "Create a new version instead",
        "New version re-runs approval and test",
        "Old versions kept for audit"]),
    ("8", "Testing & readiness", "In Test to Active", "Test team / IT", [
        "Deploy to test environment",
        "Run test quotes and orders",
        "Pass: Active, ready for CRM",
        "Fail: Rejected or new version"]),
])

# ------------------------------------------------------------------ 10 approvals
s = slide("Approval gates: nothing reaches sales by accident",
          "Each status change is a decision, recorded with who approved it and why")
gates = [("Gate 1", "Conception approved", "In Study to In Design", "Is the idea worth designing?",
          "Product Manager, Marketing, Finance"),
         ("Gate 2", "Design approved", "In Design to In Test", "Are offer, price, rules and terms right?",
          "Marketing, Finance, Legal / Regulatory, IT"),
         ("Gate 3", "Test approved", "In Test to Active", "Does it work end to end?", "Test team, IT"),
         ("Gate 4", "Start of marketing", "Active to Launched", "Go on sale on the planned date",
          "Product Manager")]
for i, (g, name, move, q, who) in enumerate(gates):
    x = Inches(0.6 + i * 3.15)
    b = box(s, x, Inches(1.8), Inches(2.9), Inches(1.1), fill=PURPLE if i < 3 else ACCENT,
            shape=MSO_SHAPE.CHEVRON if False else MSO_SHAPE.PENTAGON)
    text(s, 0, 0, 0, 0, [g, name], size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, shape=b)
    text(s, x, Inches(3.05), Inches(2.9), Inches(0.4), move, size=13, color=ACCENT, bold=True,
         align=PP_ALIGN.CENTER)
    box(s, x, Inches(3.55), Inches(2.9), Inches(2.4), fill=LILAC)
    text(s, x + Inches(0.12), Inches(3.65), Inches(2.66), Inches(2.2),
         [("Question: ", q), ("Approvers: ", who)], size=14)
text(s, Inches(0.6), Inches(6.2), Inches(12), Inches(0.6),
     "Gates and statuses: TMF620 standard. Approver roles shown are an example set; they are configured per "
     "operator.", size=12, color=GREY)

# ------------------------------------------------------------------ 11 CRM principle
s = slide("Integration with CRM: the catalog is the single source of truth",
          "CRM consumes the approved product; it never re-creates it")
cat = box(s, Inches(0.6), Inches(2.0), Inches(3.4), Inches(3.6), fill=PURPLE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 0, 0, 0, 0, ["Netcracker", "Product Catalog + PLM", "", "Only Active products are published"],
     size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=cat)
arrow(s, Inches(4.05), Inches(3.8), Inches(5.0), Inches(3.8))
pub = ["Offering and bundle structure", "Prices, discounts, contract terms",
       "Eligibility: market, channel, customer category", "Dependencies and constraints",
       "Validity dates (start / end of sale)", "Status and version"]
text(s, Inches(5.1), Inches(1.65), Inches(3.8), Inches(0.4), "What gets published", size=15, color=PURPLE,
     bold=True)
for i, p in enumerate(pub):
    b = box(s, Inches(5.1), Inches(2.1 + i * 0.6), Inches(3.8), Inches(0.5), fill=LILAC,
            shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, 0, 0, 0, 0, [p], size=13, color=DARK, anchor=MSO_ANCHOR.MIDDLE, shape=b)
arrow(s, Inches(8.95), Inches(3.8), Inches(9.9), Inches(3.8))
crm = box(s, Inches(10.0), Inches(2.0), Inches(2.8), Inches(3.6), fill=ACCENT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 0, 0, 0, 0, ["CRM", "", "Sales quote, order and contract use the catalog definition"], size=16,
     color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=crm)
text(s, Inches(0.6), Inches(5.95), Inches(12.2), Inches(0.8),
     "Netcracker: the Product Catalog supports \"sales, ordering and billing processes via industry-standard "
     "integration\"; the lead-to-order process is \"automated with a catalog-driven approach\".",
     size=13, color=GREY)

# ------------------------------------------------------------------ 12 APIs
s = slide("How the hand-over works: TM Forum Open APIs",
          "Standard interfaces between the catalog and the CRM")
table(s, Inches(0.6), Inches(1.7), Inches(12.1), [
    ["API", "Direction", "What it does", "Netcracker"],
    ["TMF620 Product Catalog", "Catalog to CRM", "CRM reads catalogs, offerings, prices, specifications", "Certified"],
    ["TMF620 events", "Catalog to CRM", "CRM is notified on create / change / status change; stays in sync", "Part of TMF620"],
    ["TMF679 Offering Qualification", "CRM to catalog", "\"Can this customer get this offer here?\"", "TM Forum standard"],
    ["TMF648 Quote Management", "Within sales", "Quotes built on catalog offerings and prices", "Certified"],
    ["TMF622 Product Ordering", "CRM to order mgmt", "Order carries offering ID and version", "Certified"],
], [3.2, 2.2, 4.9, 1.8], size=14, row_h=Inches(0.6))
steps = ["Product becomes Active", "Event sent to CRM", "CRM reads offer via TMF620",
         "Salesperson quotes; rules checked", "Order references offer version"]
for i, st in enumerate(steps):
    x = Inches(0.6 + i * 2.5)
    b = box(s, x, Inches(5.65), Inches(2.2), Inches(0.9), fill=PURPLE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, 0, 0, 0, 0, [st], size=13, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, shape=b)
    if i < 4:
        arrow(s, x + Inches(2.22), Inches(6.1), x + Inches(2.48), Inches(6.1))

# ------------------------------------------------------------------ 13 salesperson view
s = slide("What the salesperson sees", "A real Netcracker screen: the catalog drives the quote")
picture(s, "03_netcracker_real_screen_create_quote.png", Inches(0.6), Inches(1.7), Inches(7.2), Inches(4.8),
        "Source: Netcracker CPQ whitepaper (2021), Figure 3")
bullets(s, Inches(8.2), Inches(1.9), Inches(4.6), Inches(4.6), [
    ("Customer Category, Channel, Market: ", "the same eligibility rules defined in Stage 5."),
    ("Only eligible offers ", "appear for this customer."),
    ("Prices come from the catalog, ", "so the quote always uses the approved price."),
    ("The offer is validated ", "against catalog rules before the quote is submitted."),
], size=16)

# ------------------------------------------------------------------ 14 version change
s = slide("When a product changes after launch", "Versioning keeps sales and existing customers consistent")
flow = [("1", "New version created", "In the catalog, not in CRM"),
        ("2", "Approval and test", "Same gates as a new product"),
        ("3", "Change event to CRM", "New version becomes sellable"),
        ("4", "New sales use v2", "Existing customers keep v1"),
        ("5", "Old version retired", "Disappears from sales screens")]
for i, (n, t, d) in enumerate(flow):
    x = Inches(0.6 + i * 2.5)
    c = box(s, x + Inches(0.7), Inches(2.0), Inches(0.8), Inches(0.8), fill=PURPLE, shape=MSO_SHAPE.OVAL)
    text(s, 0, 0, 0, 0, [n], size=22, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, shape=c)
    if i < 4:
        arrow(s, x + Inches(1.55), Inches(2.4), x + Inches(3.15), Inches(2.4))
    text(s, x, Inches(3.0), Inches(2.2), Inches(0.5), t, size=16, color=PURPLE, bold=True, align=PP_ALIGN.CENTER)
    text(s, x, Inches(3.5), Inches(2.2), Inches(0.8), d, size=13, color=GREY, align=PP_ALIGN.CENTER)
box(s, Inches(0.6), Inches(4.7), Inches(12.1), Inches(1.5), fill=LILAC)
text(s, Inches(0.8), Inches(4.85), Inches(11.7), Inches(1.3), [
    ("Why it works: ", "every order records the offering ID and version, so delivery and billing always "
     "match what the customer bought."),
    ("Result: ", "a price or bundle change never breaks existing contracts."),
], size=16)

# ------------------------------------------------------------------ 15 Dynamics
s = slide("Example: connecting the catalog to Microsoft Dynamics 365",
          "The same pattern applies to a CRM from another vendor")
b = box(s, Inches(0.6), Inches(1.9), Inches(2.6), Inches(1.2), fill=PURPLE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 0, 0, 0, 0, ["Product catalog", "TMF620 API"], size=15, color=WHITE, bold=True,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=b)
arrow(s, Inches(3.25), Inches(2.5), Inches(4.0), Inches(2.5))
b = box(s, Inches(4.05), Inches(1.9), Inches(2.6), Inches(1.2), fill=GREY, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 0, 0, 0, 0, ["Integration layer", "maps and syncs"], size=15, color=WHITE, bold=True,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=b)
arrow(s, Inches(6.7), Inches(2.5), Inches(7.45), Inches(2.5))
b = box(s, Inches(7.5), Inches(1.9), Inches(2.6), Inches(1.2), fill=ACCENT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 0, 0, 0, 0, ["Dynamics 365", "Dataverse Web API"], size=15, color=WHITE, bold=True,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=b)
text(s, Inches(10.3), Inches(1.9), Inches(2.5), Inches(1.3),
     "No ready-made Netcracker connector; built on standard APIs", size=12, color=GREY)
table(s, Inches(0.6), Inches(3.45), Inches(12.1), [
    ["Catalog item (TMF620)", "Dynamics 365 Sales entity"],
    ["Product Offering / Bundle offering", "Product / Product bundle"],
    ["Product Catalog / Category", "Product family"],
    ["Product Offering Price / Discounts", "Price list item / Discount list"],
    ["Related offers, add-ons", "Product relationships (accessory, cross-sell, up-sell)"],
    ["Status Active / Launched / Retired", "Product status Active / Retired"],
], [6.0, 6.1], size=14, row_h=Inches(0.45))
text(s, Inches(0.6), Inches(6.35), Inches(12.1), Inches(0.5),
     "Study, design, test and approval stay in the catalog; Dynamics receives only products ready to sell.",
     size=13, color=PURPLE, bold=True)

# ------------------------------------------------------------------ 16 evidence
s = slide("Evidence: results at real operators")
for i, (op, what, res) in enumerate([
    ("Zain KSA", "Full BSS/OSS including Product Catalog", "50% decrease in product development time"),
    ("T-Mobile Netherlands", "Product Catalog + Product Lifecycle Management",
     "Products managed \"from development to production to retirement\"; fewer incorrect offers"),
    ("e& UAE", "Full BSS including product catalog", "\"Groundwork for catalog-driven offers\""),
]):
    y = Inches(1.5 + i * 1.75)
    box(s, Inches(0.6), y, Inches(7.6), Inches(1.55), fill=LILAC)
    box(s, Inches(0.6), y, Inches(0.12), Inches(1.55), fill=ACCENT if i == 0 else PURPLE)
    text(s, Inches(0.9), y + Inches(0.1), Inches(7.2), Inches(0.45), op, size=19, color=PURPLE, bold=True)
    text(s, Inches(0.9), y + Inches(0.55), Inches(7.2), Inches(0.4), what, size=13, color=GREY)
    text(s, Inches(0.9), y + Inches(0.92), Inches(7.2), Inches(0.6), res, size=15, color=DARK, bold=True)
picture(s, "04_zain_ksa_case_study_page1.png", Inches(8.6), Inches(1.5), Inches(4.2), Inches(4.9),
        "Source: Netcracker, Zain KSA case study")
text(s, Inches(0.6), Inches(6.75), Inches(8), Inches(0.3),
     "Sources: Netcracker case study and press releases with operator quotes.", size=11, color=GREY)

# ------------------------------------------------------------------ 17 close
s = slide("Summary and references")
bullets(s, Inches(0.6), Inches(1.5), Inches(6.6), Inches(5.2), [
    ("Idea to Active ", "inside the catalog: eight stages, approval gates, versions, testing."),
    ("Active to CRM ", "through standard APIs: the catalog stays the single source of truth."),
    ("Rules and prices ", "are enforced in the sales screen, not re-typed."),
    ("Changes are versioned, ", "so existing customers are never affected."),
    ("Proven ", "at Zain KSA, T-Mobile Netherlands and e&."),
], size=17)
refs = [
    "Netcracker Product Management (Product Catalog + PLM)",
    "Netcracker Commerce Management",
    "Netcracker CPQ whitepaper (2021), Figures 2 and 3",
    "Netcracker Zain KSA case study",
    "Netcracker T-Mobile Netherlands press release",
    "Netcracker e& UAE press release (Oct 2023)",
    "TM Forum SID certification report, Netcracker",
    "TM Forum Open API conformance list (June 2024)",
    "TM Forum TMF620 and TMF679 API specifications",
    "Microsoft Learn: Dynamics 365 Sales product catalog",
]
box(s, Inches(7.6), Inches(1.5), Inches(5.2), Inches(5.2), fill=LILAC)
text(s, Inches(7.8), Inches(1.6), Inches(4.8), Inches(0.4), "References (full links in the document)",
     size=14, color=PURPLE, bold=True)
text(s, Inches(7.8), Inches(2.1), Inches(4.8), Inches(4.5), [f"{i}. {r}" for i, r in enumerate(refs, 1)],
     size=12, color=DARK)

assert _page[0] == TOTAL, _page[0]
prs.save(OUT)
print("saved", OUT.name)
