"""Builds the PLC study presentation (Amdocs CatalogONE, idea to CRM).

Run:  python3 build_deck_catalogone.py
"""

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).parent
IMG = HERE / "images" / "catalogone"
OUT = HERE / "PLC_CatalogONE_Presentation.pptx"

NAVY = RGBColor(0x0C, 0x29, 0x4B)
LILAC = RGBColor(0xE8, 0xF0, 0xF8)
DARK = RGBColor(0x26, 0x22, 0x2E)
GREY = RGBColor(0x5F, 0x5A, 0x6B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0xF0, 0x73, 0x00)
MS = RGBColor(0x00, 0x78, 0xD4)
RED = RGBColor(0xB0, 0x3A, 0x2E)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
W, H = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]
TOTAL = 17
_page = [0]


def box(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    s.shadow.inherit = False
    return s


def text(slide, x, y, w, h, content, size=16, color=DARK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, shape=None):
    s = shape or slide.shapes.add_textbox(x, y, w, h)
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    items = content if isinstance(content, list) else [content]
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if isinstance(it, tuple):
            r = p.add_run()
            r.text = it[0]
            r.font.bold = True
            r.font.size = Pt(size)
            r.font.color.rgb = NAVY if color == DARK else color
            r2 = p.add_run()
            r2.text = it[1]
            r2.font.size = Pt(size)
            r2.font.color.rgb = color
        else:
            r = p.add_run()
            r.text = it
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
        if len(items) > 1:
            p.space_after = Pt(size * 0.55)
    return s


def bullets(slide, x, y, w, h, items, size=17):
    s = slide.shapes.add_textbox(x, y, w, h)
    tf = s.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(9)
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(Inches(0.3)))
        pPr.set("indent", str(-Inches(0.25)))
        pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•"}))
        lead, rest = it if isinstance(it, tuple) else ("", it)
        if lead:
            r = p.add_run()
            r.text = lead
            r.font.bold = True
            r.font.size = Pt(size)
            r.font.color.rgb = NAVY
        r = p.add_run()
        r.text = rest
        r.font.size = Pt(size)
        r.font.color.rgb = DARK
    return s


def slide(title, subtitle=None):
    _page[0] += 1
    s = prs.slides.add_slide(BLANK)
    box(s, 0, 0, W, Inches(0.12), fill=NAVY)
    text(s, Inches(0.6), Inches(0.32), Inches(12), Inches(0.7), title, size=28, color=NAVY, bold=True)
    if subtitle:
        text(s, Inches(0.6), Inches(1.0), Inches(12), Inches(0.45), subtitle, size=15, color=GREY)
    text(s, Inches(0.6), Inches(7.05), Inches(9.5), Inches(0.3),
         "PLC study  |  Amdocs CatalogONE: from idea to CRM", size=10, color=GREY)
    text(s, Inches(11.7), Inches(7.05), Inches(1.0), Inches(0.3), f"{_page[0]} / {TOTAL}", size=10,
         color=GREY, align=PP_ALIGN.RIGHT)
    return s


def picture(s, file, x, y, max_w, max_h, caption=None):
    iw, ih = Image.open(IMG / file).size
    scale = min(max_w / iw, max_h / ih)
    w, h = int(iw * scale), int(ih * scale)
    px = x + (max_w - w) // 2
    box(s, px - Inches(0.04), y - Inches(0.04), w + Inches(0.08), h + Inches(0.08),
        line=RGBColor(0xC5, 0xD4, 0xE4))
    s.shapes.add_picture(str(IMG / file), px, y, w, h)
    if caption:
        text(s, x, y + h + Inches(0.06), max_w, Inches(0.45), caption, size=10, color=GREY,
             align=PP_ALIGN.CENTER)


def arrow(s, x1, y1, x2, y2, color=NAVY):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.line.color.rgb = color
    c.line.width = Pt(2)
    ln = c.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle"}))
    return c


def table(s, x, y, w, rows, col_w, size=13, row_h=Inches(0.42)):
    shp = s.shapes.add_table(len(rows), len(rows[0]), x, y, w, row_h * len(rows))
    t = shp.table
    for c, cw in enumerate(col_w):
        t.columns[c].width = Inches(cw)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            cell.text = val
            cell.margin_left = cell.margin_right = Inches(0.08)
            p = cell.text_frame.paragraphs[0]
            p.runs[0].font.size = Pt(size)
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = NAVY
                p.runs[0].font.bold = True
                p.runs[0].font.color.rgb = WHITE
            else:
                cell.fill.fore_color.rgb = LILAC if r % 2 else WHITE
                p.runs[0].font.color.rgb = DARK
                p.runs[0].font.bold = c == 0
    return t


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t


# ------------------------------------------------------------------ 1 title
_page[0] += 1
s = prs.slides.add_slide(BLANK)
box(s, 0, 0, W, H, fill=NAVY)
box(s, Inches(0.8), Inches(2.1), Inches(0.12), Inches(2.6), fill=ACCENT)
text(s, Inches(1.1), Inches(2.0), Inches(11), Inches(0.5), "PLC STUDY", size=16, color=LILAC, bold=True)
text(s, Inches(1.1), Inches(2.5), Inches(11.5), Inches(1.5),
     "Product Lifecycle Management in Amdocs CatalogONE", size=36, color=WHITE, bold=True)
text(s, Inches(1.1), Inches(4.1), Inches(11.5), Inches(0.9),
     "A catalog-only PLC that publishes to consuming applications — including the customer's CRM",
     size=20, color=LILAC)
text(s, Inches(1.1), Inches(5.7), Inches(11), Inches(0.45),
     "Prepared by: [Your name]   |   October 2026   |   Draft 4", size=14, color=LILAC)
notes(s, "CatalogONE is the PLC. It publishes to consuming apps and 'the customer's CRM'. "
         "Dynamics 365 is named on the Amdocs–Microsoft CEP, not as a CatalogONE connector.")

# ------------------------------------------------------------------ 2 agenda
s = slide("Agenda")
items = [
    ("01", "Why this PLC", "Catalog only; CRM is not replaced"),
    ("02", "PLC and TM Forum", "The industry lifecycle model"),
    ("03", "CatalogONE at a glance", "Modules and data model"),
    ("04", "Using the system", "Eight stages from idea to launch-ready"),
    ("05", "How catalog reaches CRM", "CatalogONE vs Customer Engagement Platform"),
    ("06", "Evidence", "Results at real operators"),
]
for i, (n, t, d) in enumerate(items):
    col, row = i % 2, i // 2
    x, y = Inches(0.8 + col * 6.1), Inches(1.55 + row * 1.7)
    box(s, x, y, Inches(5.7), Inches(1.4), fill=LILAC)
    text(s, x + Inches(0.2), y + Inches(0.2), Inches(1.1), Inches(1.0), n, size=36, color=NAVY, bold=True)
    text(s, x + Inches(1.3), y + Inches(0.25), Inches(4.3), Inches(0.5), t, size=20, color=DARK, bold=True)
    text(s, x + Inches(1.3), y + Inches(0.75), Inches(4.3), Inches(0.5), d, size=14, color=GREY)

# ------------------------------------------------------------------ 3 takeaways
s = slide("Why this PLC: catalog only. CRM is not replaced.",
          "Two facts, kept separate: CatalogONE publishes to CRM. The Customer Engagement Platform names Dynamics 365.")
cards = [
    ("Catalog, not CRM", "CatalogONE owns products, prices, rules and approvals. "
     "It publishes to consuming apps and \"the customer's CRM\"."),
    ("One source of truth", "Commercial and technical definitions live once. "
     "CRM, ordering and billing consume them."),
    ("A governed journey", "Ideation to end-of-life, with workflow, approvals, "
     "simulations and versioned change."),
    ("Business users configure", "Role-based UI, templates and building blocks. "
     "No code. No deployment overhead."),
    ("Publish, do not re-key", "Kafka + TMF620 to consuming applications. "
     "Dynamics 365 is named on the Amdocs–Microsoft platform, not on CatalogONE."),
    ("Proven results", "European operator: 9 months to 1–2 weeks. "
     "Sprint / T-Mobile: expected 50% faster TTM."),
]
for i, (t, d) in enumerate(cards):
    col, row = i % 3, i // 3
    x, y = Inches(0.6 + col * 4.1), Inches(1.6 + row * 2.55)
    box(s, x, y, Inches(3.8), Inches(2.3), fill=LILAC)
    box(s, x, y, Inches(3.8), Inches(0.1), fill=MS if i == 4 else (ACCENT if i == 5 else NAVY))
    text(s, x + Inches(0.15), y + Inches(0.25), Inches(3.5), Inches(0.5), t, size=17, color=NAVY, bold=True)
    text(s, x + Inches(0.15), y + Inches(0.8), Inches(3.5), Inches(1.35), d, size=13, color=DARK)
notes(s, "Do not say CatalogONE sends data to Dynamics. Say: publishes to CRM; CEP names Dynamics.")

# ------------------------------------------------------------------ 4 TMF620 lifecycle
s = slide("PLC and the TM Forum standard in brief",
          "TMF620 defines the statuses. CatalogONE PLM runs them. CRM sees Active / Launched after publish.")
stages = [("In Study", "Idea studied"), ("In Design", "Being designed"), ("In Test", "Being tested"),
          ("Active", "Ready, not yet sold"), ("Launched", "On sale"), ("Retired", "No new sales")]
gates = ["Conception\napproved", "Design\napproved", "Test\napproved", "Start of\nmarketing", "End of\nmarketing"]
bw, gap = Inches(1.75), Inches(0.42)
x0, y0 = Inches(0.6), Inches(2.55)
for i, (st, d) in enumerate(stages):
    x = x0 + i * (bw + gap)
    fill = NAVY if i < 4 else (MS if i == 4 else GREY)
    b = box(s, x, y0, bw, Inches(1.1), fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, 0, 0, 0, 0, [st], size=18, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, shape=b)
    text(s, x, y0 + Inches(1.2), bw, Inches(0.5), d, size=13, color=GREY, align=PP_ALIGN.CENTER)
    if i < 5:
        arrow(s, x + bw + Inches(0.03), y0 + Inches(0.55), x + bw + gap - Inches(0.03), y0 + Inches(0.55))
        text(s, x + bw - Inches(0.45), y0 - Inches(0.85), gap + Inches(0.9), Inches(0.7), gates[i], size=11,
             color=NAVY, align=PP_ALIGN.CENTER)
rx = x0 + 2 * (bw + gap)
b = box(s, rx, Inches(4.5), bw, Inches(0.55), fill=RED, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 0, 0, 0, 0, ["Rejected (test KO)"], size=13, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
     anchor=MSO_ANCHOR.MIDDLE, shape=b)
box(s, x0 - Inches(0.1), Inches(5.45), 4 * (bw + gap) - gap + Inches(0.2), Inches(0.08), fill=NAVY)
text(s, x0, Inches(5.6), Inches(8), Inches(0.4), "Inside CatalogONE: defined, approved and tested",
     size=14, color=NAVY, bold=True)
text(s, x0 + 4 * (bw + gap), Inches(5.6), Inches(4.2), Inches(0.4), "Sales CRM: visible after publish",
     size=14, color=MS, bold=True)
text(s, x0, Inches(6.2), Inches(12), Inches(0.5),
     "Amdocs is TM Forum-certified on TMF620. CatalogONE APIs use the TMF data model.",
     size=13, color=GREY)

# ------------------------------------------------------------------ 5 modules
s = slide("CatalogONE at a glance",
          "Standalone telecom catalog + PLM. Modular. Does not replace CRM.")
picture(s, "09_modules_list.png", Inches(0.5), Inches(1.55), Inches(6.0), Inches(4.85),
        "Source: Amdocs, 6 Steps ebook, p.19")
table(s, Inches(6.7), Inches(1.55), Inches(6.1), [
    ["Module", "Role"],
    ["PLM", "Ideation to end-of-life; approvals"],
    ["Commerce", "Offers, bundles, channels, segments"],
    ["Centralized", "Master catalog; publish downstream"],
    ["Service / resource", "Technical specs for delivery"],
    ["Templates", "Reusable telco building blocks"],
    ["Federation", "Partner catalogs via TMF APIs"],
    ["iCatalog", "KPI dashboards and recommendations"],
    ["Charging (optional)", "Rated usage — not required here"],
], [2.4, 3.7], size=13, row_h=Inches(0.52))

# ------------------------------------------------------------------ 6 catalog / CRM split
s = slide("Where Dynamics 365 is named: Customer Engagement Platform (CEP)",
          "CEP = Amdocs catalog + CPQ + commerce + OM, on Dynamics 365. Not a CatalogONE plug-in.")
picture(s, "01_cep_architecture.png", Inches(0.5), Inches(1.5), Inches(7.6), Inches(4.7),
        "Source: Amdocs and Microsoft CEP paper, p.3 — this is CEP, not a CatalogONE datasheet")
for i, (t, d) in enumerate([
    ("Catalog (Amdocs)", "Product lifecycle, offerings, prices, rules. Source of truth."),
    ("Sales / Service (Dynamics 365)", "In CEP, quotes, orders and cases sit here. CRM is not replaced."),
    ("Live operator: PLDT", "IDC 2026: case management on CEP. Agents use Dynamics 365. Care, not a CatalogONE catalog feed."),
]):
    y = Inches(1.55 + i * 1.55)
    box(s, Inches(8.35), y, Inches(4.45), Inches(1.4), fill=LILAC)
    text(s, Inches(8.5), y + Inches(0.12), Inches(4.15), Inches(0.4), t, size=16, color=NAVY, bold=True)
    text(s, Inches(8.5), y + Inches(0.52), Inches(4.15), Inches(0.8), d, size=13, color=DARK)

# ------------------------------------------------------------------ 7 data model
s = slide("Core data model, with one running example",
          "TM Forum entities in CatalogONE, illustrated with \"Business Internet 100\"")
model = [("Product Specification", "What it technically is", "Dedicated Internet Access, 100 Mbps"),
         ("Product Offering", "What is sold (single or bundle)", "\"Business Internet 100\" bundle"),
         ("Product Offering Price", "Prices attached to the offer", "Monthly fee + installation per site"),
         ("Channel / segment", "Where and to whom it is sold", "Enterprise AMs; SME and corporate"),
         ("Product", "What the customer owns after sale", "Company X, 3 sites")]
for i, (e, m, ex) in enumerate(model):
    y = Inches(1.6 + i * 0.9)
    b = box(s, Inches(0.8), y, Inches(3.6), Inches(0.75), fill=NAVY if i < 4 else GREY,
            shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, 0, 0, 0, 0, [e], size=15, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, shape=b)
    text(s, Inches(4.6), y + Inches(0.15), Inches(4.0), Inches(0.5), m, size=15, color=DARK)
    b2 = box(s, Inches(8.6), y, Inches(4.1), Inches(0.75), fill=LILAC, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, 0, 0, 0, 0, [ex], size=14, color=NAVY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=b2)
text(s, Inches(8.6), Inches(1.25), Inches(4.1), Inches(0.35), "Example (illustrative)", size=12, color=GREY,
     align=PP_ALIGN.CENTER)
text(s, Inches(0.8), Inches(6.2), Inches(12), Inches(0.55),
     "\"Business Internet 100\" = 100 Mbps dedicated internet + static IPs + managed router + SLA, "
     "on 12 / 24 / 36-month contracts. Unrated enterprise product — Charging catalog is not required.",
     size=13, color=GREY)


def stage_slide(title, subtitle, cards):
    s = slide(title, subtitle)
    for i, (num, name, status, owner, steps) in enumerate(cards):
        x = Inches(0.55 + i * 3.15)
        y = Inches(1.55)
        box(s, x, y, Inches(3.0), Inches(5.25), fill=LILAC)
        hb = box(s, x, y, Inches(3.0), Inches(1.0), fill=NAVY)
        text(s, 0, 0, 0, 0, [f"Stage {num}", name], size=15, color=WHITE, bold=True, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.MIDDLE, shape=hb)
        text(s, x + Inches(0.1), y + Inches(1.08), Inches(2.8), Inches(0.32), status, size=12, color=ACCENT,
             bold=True)
        text(s, x + Inches(0.1), y + Inches(1.38), Inches(2.8), Inches(0.4), "Owner: " + owner, size=12,
             color=GREY)
        bullets(s, x + Inches(0.02), y + Inches(1.85), Inches(2.95), Inches(3.2), steps, size=12)
    return s


s = stage_slide("Using the system (1/2): define the product",
                "What the user does, stage by stage (screen names are illustrative)", [
    ("1", "Ideation & business case", "In Study", "Product Manager", [
        "Register \"Business Internet 100 for SMEs\"",
        "Attach business case: segment, sites, revenue",
        "Use iCatalog to size demand and cannibalisation",
        "Reuse existing catalog items",
        "Submit for conception approval"]),
    ("2", "Technical definition", "In Design", "Product Specialist", [
        "Open the service / resource catalog",
        "Set bandwidth, access type, IP block",
        "Add router, IP block, SLA specs",
        "Expose them to the commercial catalog"]),
    ("3", "Commercial offer", "In Design", "Product Manager", [
        "Start the Commerce offer wizard",
        "Bundle internet + IPs + router + SLA",
        "Add optional add-ons",
        "Place in the Enterprise catalog"]),
    ("4", "Pricing", "In Design", "PM + Finance", [
        "Monthly fee and installation per site",
        "Prices by contract term (12 / 24 / 36)",
        "Multi-site discount, validity dates",
        "Charging catalog is optional here"]),
])
notes(s, "Evidence: role-based UI; Commerce wizard; iCatalog for ideation.")

s = stage_slide("Using the system (2/2): control and release it",
                "Rules, approvals, versions and testing decide when a product is ready", [
    ("5", "Rules & eligibility", "In Design", "Product Manager", [
        "Market: fibre-covered areas",
        "Channel: enterprise account managers",
        "Segment: SME, corporate",
        "Constraints (backup link only on 24/36 months)"]),
    ("6", "Approval gates", "In Design to In Test", "Approvers", [
        "PM submits in PLM",
        "Marketing, Finance, Legal, IT review",
        "Approve or reject with comments",
        "All approved: status moves to In Test"]),
    ("7", "Versioning", "Any status", "PM / Catalog admin", [
        "Never edit a launched product directly",
        "Create a new version instead",
        "New version re-runs approval and test",
        "Auto-release publishes the new version"]),
    ("8", "Testing & readiness", "In Test to Active", "Test team / IT", [
        "Release to sandbox",
        "Test quotes and orders",
        "Pass: Active, ready to publish",
        "Fail: Rejected or new version"]),
])

# ------------------------------------------------------------------ 10 approvals
s = slide("Approval gates: nothing is published by accident",
          "Each status change is a decision, recorded with who approved it and why")
gates = [("Gate 1", "Conception approved", "In Study to In Design", "Is the idea worth designing?",
          "Product Manager, Marketing, Finance"),
         ("Gate 2", "Design approved", "In Design to In Test", "Are offer, price, rules and terms right?",
          "Marketing, Finance, Legal / Regulatory, IT"),
         ("Gate 3", "Test approved", "In Test to Active", "Does it work end to end?", "Test team, IT"),
         ("Gate 4", "Start of marketing", "Active to Launched", "Go on sale on the planned date",
          "Product Manager")]
for i, (g, name, move, q, who) in enumerate(gates):
    x = Inches(0.6 + i * 3.15)
    b = box(s, x, Inches(1.7), Inches(2.95), Inches(1.1), fill=NAVY if i < 3 else MS,
            shape=MSO_SHAPE.PENTAGON)
    text(s, 0, 0, 0, 0, [g, name], size=15, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, shape=b)
    text(s, x, Inches(2.95), Inches(2.95), Inches(0.4), move, size=13, color=ACCENT, bold=True,
         align=PP_ALIGN.CENTER)
    box(s, x, Inches(3.45), Inches(2.95), Inches(2.4), fill=LILAC)
    text(s, x + Inches(0.12), Inches(3.55), Inches(2.7), Inches(2.2),
         [("Question: ", q), ("Approvers: ", who)], size=14)
text(s, Inches(0.6), Inches(6.1), Inches(12), Inches(0.6),
     "Gates and statuses: TMF620. CatalogONE PLM supplies workflow, approvals and notifications "
     "(KT: business users manage the offering lifecycle).", size=12, color=GREY)

# ------------------------------------------------------------------ 11 CRM principle
s = slide("The catalog publishes. CRM consumes. Two documented facts.",
          "CatalogONE names CRM. CEP names Dynamics 365. They are not the same sentence.")
cat = box(s, Inches(0.55), Inches(1.85), Inches(3.5), Inches(3.7), fill=NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 0, 0, 0, 0, ["CatalogONE", "Commerce + PLM", "", "Publishes Active products to consuming applications"],
     size=15, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=cat)
arrow(s, Inches(4.1), Inches(3.7), Inches(5.05), Inches(3.7))
pub = ["Offering and bundle structure", "Prices, discounts, contract terms",
       "Eligibility: channel, segment, market", "Dependencies and constraints",
       "Validity dates (start / end of sale)", "Status and version"]
text(s, Inches(5.15), Inches(1.5), Inches(3.9), Inches(0.35), "What gets published", size=14, color=NAVY,
     bold=True)
for i, p in enumerate(pub):
    b = box(s, Inches(5.15), Inches(1.95 + i * 0.58), Inches(3.9), Inches(0.5), fill=LILAC,
            shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, 0, 0, 0, 0, [p], size=13, color=DARK, anchor=MSO_ANCHOR.MIDDLE, shape=b)
arrow(s, Inches(9.1), Inches(3.7), Inches(10.05), Inches(3.7), color=MS)
crm = box(s, Inches(10.15), Inches(1.85), Inches(2.7), Inches(3.7), fill=MS, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 0, 0, 0, 0, ["Customer's CRM", "(unnamed on CatalogONE)", "", "Dynamics 365 only where CEP is the platform"],
     size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=crm)
text(s, Inches(0.55), Inches(5.8), Inches(12.2), Inches(0.9),
     "Fact 1 — CatalogONE: \"publishing catalog information to multiple consuming applications\" and "
     "\"integration to … the customer's CRM solution\". Fact 2 — CEP: unified Catalog plus "
     "\"Dynamics 365 modules focused on sales, marketing, and customer service\".",
     size=13, color=GREY)

# ------------------------------------------------------------------ 12 APIs
s = slide("How CatalogONE publishes — destination is CRM, not named as Dynamics",
          "Master-catalog publish + Kafka events + TM Forum Open APIs")
table(s, Inches(0.55), Inches(1.5), Inches(12.2), [
    ["Mechanism", "Direction", "What it does"],
    ["CatalogONE Centralized", "Catalog → consuming apps", "Author once; publish offerings, prices, rules"],
    ["Kafka distribution", "Catalog → downstream systems", "Sync approved versions in minutes, not weeks"],
    ["TMF620 Product Catalog", "Catalog → any TMF consumer", "Read catalogs, offerings, prices, specifications"],
    ["TMF620 events", "Catalog → subscribers", "Notify on create / change / status change"],
    ["TMF679 Qualification", "CRM / CPQ → catalog", "\"Can this customer get this offer here?\""],
    ["TMF622 Product Ordering", "CRM / CPQ → order mgmt", "Order carries offering ID and version"],
], [3.3, 3.2, 5.7], size=13, row_h=Inches(0.5))
steps = ["Product becomes Active", "Kafka / TMF event", "CRM reads TMF620",
         "AM quotes in CRM", "Order stores offer version"]
for i, st in enumerate(steps):
    x = Inches(0.55 + i * 2.5)
    b = box(s, x, Inches(5.55), Inches(2.25), Inches(0.95), fill=NAVY if i < 4 else MS,
            shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, 0, 0, 0, 0, [st], size=13, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, shape=b)
    if i < 4:
        arrow(s, x + Inches(2.28), Inches(6.02), x + Inches(2.48), Inches(6.02),
              color=MS if i == 3 else NAVY)

# ------------------------------------------------------------------ 13 complementary stack
s = slide("What sits around the catalog — on CEP, not on a CatalogONE datasheet",
          "Digital Enterprise Catalog (Amdocs) and Sales management (Dynamics 365) on one CEP platform")
picture(s, "02_cep_catalog_and_crm.png", Inches(0.5), Inches(1.5), Inches(12.3), Inches(4.9),
        "Source: Amdocs and Microsoft Customer Engagement Platform paper, p.4")

# ------------------------------------------------------------------ 14 version change
s = slide("When a product changes after launch", "Versioning keeps sales and existing customers consistent")
flow = [("1", "New version created", "In CatalogONE, not in CRM"),
        ("2", "Approval and test", "Same gates as a new product"),
        ("3", "Event to CRM", "New version becomes sellable"),
        ("4", "New sales use v2", "Existing customers keep v1"),
        ("5", "Old version retired", "Disappears from sales screens")]
for i, (n, t, d) in enumerate(flow):
    x = Inches(0.6 + i * 2.5)
    c = box(s, x + Inches(0.7), Inches(1.85), Inches(0.8), Inches(0.8), fill=NAVY, shape=MSO_SHAPE.OVAL)
    text(s, 0, 0, 0, 0, [n], size=22, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, shape=c)
    if i < 4:
        arrow(s, x + Inches(1.55), Inches(2.25), x + Inches(3.15), Inches(2.25))
    text(s, x, Inches(2.85), Inches(2.2), Inches(0.5), t, size=16, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    text(s, x, Inches(3.35), Inches(2.2), Inches(0.8), d, size=13, color=GREY, align=PP_ALIGN.CENTER)
box(s, Inches(0.6), Inches(4.45), Inches(12.1), Inches(2.15), fill=LILAC)
text(s, Inches(0.8), Inches(4.6), Inches(11.7), Inches(1.9), [
    ("Why it works: ", "every order records the offering ID and version, so delivery and billing always "
     "match what the customer bought."),
    ("If the CRM is Dynamics 365 [E]: ", "new version maps to Under Revision, then Active. Retired "
     "catalog versions are Retired and drop off the quote screen."),
    ("Result: ", "a price or bundle change never breaks existing contracts."),
], size=15)

# ------------------------------------------------------------------ 15 Dynamics mapping
s = slide("Example only: TMF620 items on Dynamics 365 Sales [E]",
          "Not a vendor connector spec. Same pattern as any TMF620 catalog, if Dynamics is the CRM.")
b = box(s, Inches(0.55), Inches(1.5), Inches(2.7), Inches(1.05), fill=NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 0, 0, 0, 0, ["CatalogONE", "TMF620 + Kafka"], size=15, color=WHITE, bold=True,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=b)
arrow(s, Inches(3.3), Inches(2.02), Inches(4.05), Inches(2.02))
b = box(s, Inches(4.1), Inches(1.5), Inches(2.7), Inches(1.05), fill=GREY, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 0, 0, 0, 0, ["Integration", "maps into Dataverse"], size=15, color=WHITE, bold=True,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=b)
arrow(s, Inches(6.85), Inches(2.02), Inches(7.6), Inches(2.02), color=MS)
b = box(s, Inches(7.65), Inches(1.5), Inches(2.7), Inches(1.05), fill=MS, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 0, 0, 0, 0, ["Dynamics 365", "product catalog"], size=15, color=WHITE, bold=True,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=b)
text(s, Inches(10.5), Inches(1.5), Inches(2.3), Inches(1.15),
     "Pre-built Catalog+Dynamics is the Customer Engagement Platform. Not a CatalogONE plug-in.",
     size=12, color=GREY)
table(s, Inches(0.55), Inches(2.85), Inches(12.2), [
    ["Catalog / TMF620 item", "Dynamics 365 Sales entity [M] — example [E]"],
    ["Product Offering / Bundle offering", "Product / Product bundle"],
    ["Product Catalog / Category", "Product family"],
    ["Product Offering Price / Discounts", "Price list item / Discount list"],
    ["Related offers, add-ons", "Product relationships (accessory, cross-sell, up-sell)"],
    ["Status Active / Launched / Retired", "Product status Active / Retired"],
], [6.1, 6.1], size=14, row_h=Inches(0.48))
text(s, Inches(0.55), Inches(6.25), Inches(12.2), Inches(0.5),
     "Two paths: buy the Customer Engagement Platform, or keep Dynamics and integrate CatalogONE via TMF620 / Kafka.",
     size=13, color=NAVY, bold=True)

# ------------------------------------------------------------------ 16 evidence
s = slide("Evidence: results at real operators",
          "CatalogONE cases first. PLDT is CEP customer care with Dynamics 365 — not a CatalogONE catalog feed.")
for i, (op, what, res) in enumerate([
    ("European quad-play", "CatalogONE unifying siloed LoBs",
     "New offers: 9 months → 1–2 weeks. Changes in hours. 80% by configuration."),
    ("Sprint / T-Mobile US", "CatalogONE + DigitalONE",
     "Expected 50% TTM cut. 22 channel catalogs → one. Business users define offers."),
    ("KT Corporation", "CatalogONE on KT private cloud",
     "Business-driven UI; foundation to accelerate 5G consumer and enterprise offers."),
    ("PLDT (Philippines)", "CEP Unified Case Management — agents use Dynamics 365",
     "IDC: 6 weeks; 48% faster L1; +27 NPS. Care, not catalog publish."),
]):
    y = Inches(1.5 + i * 1.22)
    box(s, Inches(0.55), y, Inches(7.5), Inches(1.12), fill=LILAC)
    box(s, Inches(0.55), y, Inches(0.12), Inches(1.12), fill=MS if i == 3 else (ACCENT if i == 0 else NAVY))
    text(s, Inches(0.85), y + Inches(0.06), Inches(7.05), Inches(0.32), op, size=16, color=NAVY, bold=True)
    text(s, Inches(0.85), y + Inches(0.36), Inches(7.05), Inches(0.26), what, size=12, color=GREY)
    text(s, Inches(0.85), y + Inches(0.66), Inches(7.05), Inches(0.38), res, size=13, color=DARK, bold=True)
picture(s, "10_sprint_results.png", Inches(8.25), Inches(1.5), Inches(4.55), Inches(5.0),
        "Source: TM Forum, Sprint case study")

# ------------------------------------------------------------------ 17 close
s = slide("Summary and references")
bullets(s, Inches(0.55), Inches(1.4), Inches(6.7), Inches(5.3), [
    ("Idea to Active ", "inside CatalogONE: eight stages, approval gates, versions, testing."),
    ("CatalogONE publishes ", "to consuming applications and \"the customer's CRM\" "
     "via Centralized, Kafka and TMF620."),
    ("Dynamics 365 is named on the Customer Engagement Platform, ",
     "as Sales and Customer Service next to Catalog — not on the CatalogONE datasheet. "
     "PLDT is live on that platform for care."),
    ("Two paths: ", "buy the Customer Engagement Platform, or keep Dynamics and "
     "integrate CatalogONE via TMF620. CRM is not replaced."),
    ("Changes are versioned, ", "so existing customers are never affected."),
    ("Proven CatalogONE ", "at a European quad-play, Sprint / T-Mobile and KT. "
     "Customer Engagement Platform + Dynamics 365 is live at PLDT for case management."),
], size=14)
refs = [
    "Amdocs CatalogONE datasheets (2025)",
    "Amdocs, 6 Steps catalog ebook (2025)",
    "Amdocs CatalogONE whitepaper (role-based UI)",
    "Appledore Research, CatalogONE profile (2020)",
    "Amdocs–Microsoft CEP paper and datasheet",
    "Microsoft AppSource: Customer Engagement Platform",
    "IDC PLDT case study (CEP Unified Case Management, 2026)",
    "TM Forum Sprint / T-Mobile case study",
    "Amdocs KT CatalogONE press release (Nov 2020)",
    "TM Forum Open API conformance (Amdocs TMF620)",
    "Microsoft Learn: Dynamics 365 product catalog",
]
box(s, Inches(7.5), Inches(1.4), Inches(5.3), Inches(5.3), fill=LILAC)
text(s, Inches(7.7), Inches(1.5), Inches(4.95), Inches(0.4), "References (full links in the document)",
     size=14, color=NAVY, bold=True)
text(s, Inches(7.7), Inches(2.0), Inches(4.95), Inches(4.5), [f"{i}. {r}" for i, r in enumerate(refs, 1)],
     size=12, color=DARK)

assert _page[0] == TOTAL, _page[0]
prs.save(OUT)
print("saved", OUT.name)
