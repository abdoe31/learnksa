"""Builds the PLC study document (Netcracker Product Management).

Run:  pip install python-docx  &&  python3 build_document.py
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).parent
IMG = HERE / "images"
OUT = HERE / "PLC_Netcracker_Study.docx"

PURPLE = RGBColor(0x5B, 0x2C, 0x83)
GREY = RGBColor(0x5F, 0x5A, 0x6B)

doc = Document()
for section in doc.sections:
    section.left_margin = section.right_margin = Cm(2.2)
    section.top_margin = section.bottom_margin = Cm(2.0)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
    st = doc.styles[name]
    st.font.name = "Calibri"
    st.font.size = Pt(size)
    st.font.color.rgb = PURPLE


def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def h1(t):
    doc.add_heading(t, level=1)


def h2(t):
    doc.add_heading(t, level=2)


def h3(t):
    doc.add_heading(t, level=3)


def para(t, italic=False, color=None, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
    r = p.add_run(t)
    r.italic = italic
    if color:
        r.font.color.rgb = color
    return p


def bullets(items, style="List Bullet"):
    for it in items:
        p = doc.add_paragraph(style=style)
        if isinstance(it, tuple):
            r = p.add_run(it[0])
            r.bold = True
            p.add_run(it[1])
        else:
            p.add_run(it)


def numbered(items):
    for i, it in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(f"{i}.  ")
        r.bold = True
        r.font.color.rgb = PURPLE
        p.add_run(it)


def table(rows, widths=None, header_fill="5B2C83"):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            cell.text = ""
            run = cell.paragraphs[0].add_run(val)
            run.font.size = Pt(9.5)
            if r == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                shade(cell, header_fill)
            elif c == 0:
                run.bold = True
            if widths:
                cell.width = Cm(widths[c])
            if r < len(rows) - 1 and len(rows) <= 12:
                cell.paragraphs[0].paragraph_format.keep_with_next = True
    if widths:
        t.autofit = False
        for c, w in enumerate(widths):
            t.columns[c].width = Cm(w)
    doc.add_paragraph()
    return t


def note(title, text):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    cell = t.cell(0, 0)
    shade(cell, "EEE7F5")
    p = cell.paragraphs[0]
    r = p.add_run(title + "  ")
    r.bold = True
    r.font.color.rgb = PURPLE
    p.add_run(text).font.size = Pt(10)
    doc.add_paragraph()


def image(file, caption, width_cm=16):
    doc.add_picture(str(IMG / file), width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(caption)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = GREY


def stage(title, status, purpose, who, inputs, steps, outputs, evidence, gaps=None):
    h2(title)
    table([
        ["Item", "Detail"],
        ["Lifecycle status", status],
        ["Purpose", purpose],
        ["Who (proposed roles) [P]", who],
        ["Inputs", inputs],
        ["Exit criteria / output", outputs],
    ], widths=[4.2, 12.4])
    h3("What you do in the system (walkthrough)")
    numbered(steps)
    h3("Evidence")
    bullets(evidence)
    if gaps:
        h3("Not publicly documented (confirm in a vendor demo)")
        bullets(gaps)


# ------------------------------------------------------------------ cover
t = doc.add_paragraph()
r = t.add_run("PLC STUDY  |  WE DATA")
r.bold = True
r.font.color.rgb = PURPLE
r.font.size = Pt(11)
title = doc.add_paragraph()
r = title.add_run("Product Lifecycle Management in Netcracker:\nfrom Idea to CRM")
r.bold = True
r.font.size = Pt(22)
para("System studied: Netcracker Product Management (Product Catalog + Product Lifecycle Management), "
     "part of Netcracker Digital BSS / Commerce Management. Vendor: Netcracker Technology, "
     "a wholly owned subsidiary of NEC Corporation.", color=GREY)
para("Prepared by: [Your name]    |    Version: Draft 2 for review    |    October 2026", color=GREY)

note("How to read this document.",
     "Every statement is tagged with its source: [N] = documented by Netcracker; "
     "[S] = defined by TM Forum standards (TMF620 and related Open APIs); "
     "[M] = Microsoft documentation; [E] = illustrative example. "
     "Netcracker's detailed user guides are not public, so screen names in the walkthroughs are "
     "illustrative and will differ in a real installation.")

# ------------------------------------------------------------------ 0
h1("Key takeaways")
bullets([
    ("One source of truth. ", "Every product, price and rule is defined once in the catalog; CRM, ordering "
     "and billing read it instead of keeping their own copies [N]."),
    ("A governed journey. ", "A product moves through fixed statuses (In Study, In Design, In Test, Active) "
     "and cannot reach sales until each approval gate is passed [S]."),
    ("Business users configure, not IT. ", "Offers are built from reusable templates and components, so new "
     "products need configuration rather than development [N]."),
    ("Rules travel with the product. ", "Eligibility by market, channel and customer category is defined in "
     "the catalog and enforced automatically when a salesperson quotes in CRM [N]."),
    ("Standard integration. ", "The catalog is exposed to CRM through TM Forum Open APIs that Netcracker is "
     "certified on, so any CRM can consume it [N][S]."),
    ("Proven results. ", "Zain KSA reported a 50% decrease in product development time; T-Mobile Netherlands "
     "manages products \"from development to production to retirement\" in one place [N]."),
])

# ------------------------------------------------------------------ 1
doc.add_page_break()
h1("1. PLC and the TM Forum standard in brief")
para("A Product Lifecycle Management (PLC) system is where an operator turns an idea into a sellable "
     "product, controls every change to it, and publishes it to the systems that sell, deliver and bill it.")
para("TM Forum, the telecom industry standards body, defines this lifecycle in its Product Catalog "
     "Management API (TMF620). Each catalog item carries one of these statuses [S]:")
table([
    ["Status", "Meaning", "Gate to enter it"],
    ["In Study", "The idea is being studied", "Idea registered"],
    ["In Design", "The product is being designed", "Conception approved"],
    ["In Test", "The product is being tested", "Design approved"],
    ["Rejected", "Test failed (final status)", "Test KO"],
    ["Active", "Ready, but not yet sold", "Test approved"],
    ["Launched", "Available to customers through sales channels", "Start of marketing"],
    ["Retired / Obsolete", "No new sales / no customers left", "End of marketing"],
], widths=[3.2, 8.2, 5.2])
para("Netcracker is certified by TM Forum on TMF620 (conformance list, June 2024), so the lifecycle described "
     "below follows the industry model rather than a vendor-specific one [N]. This document covers everything "
     "up to Launched: the stages before CRM, and the hand-over to CRM.", italic=True)

# ------------------------------------------------------------------ 2
h1("2. Netcracker Product Management at a glance")
para("Netcracker groups product capabilities under the Product Management domain, with two core "
     "components [N]:")
bullets([
    ("Product Catalog: ", "\"a unified holistic view and easy access to up-to-date product information "
     "(pricing, discounts, fees, loyalty programs, taxes), supporting sales, ordering and billing processes "
     "via industry-standard integration.\" [N]"),
    ("Product Lifecycle Management: ", "\"end-to-end lifecycle management from marketing concept development "
     "and product configuration to product testing, test launch and subsequent mass production. It also "
     "fully supports product versioning, upgrades and retirement.\" [N]"),
])
image("01_tmforum_netcracker_module_map.png",
      "Figure 1. Netcracker module map from TM Forum's certification report. Product Management = Product "
      "Information Management (catalog) + Product Lifecycle Management. Source: TM Forum SID certification "
      "report, Netcracker TOMS Release 8.2.")

h2("2.1 Modules")
table([
    ["Module", "What it does", "Source"],
    ["Product Lifecycle Management", "Drives the product through concept, configuration, testing and launch; "
     "manages versions and approvals", "[N]"],
    ["Product Offering Catalog", "The commercial side: offers, bundles, business rules, prices, discounts", "[N]"],
    ["Technical Product Catalog", "Product and service specifications, provisioning templates, and "
     "decomposition rules (how an order splits into technical tasks)", "[N]"],
    ["Templates & reusable components", "Pre-defined product templates and out-of-the-box components for "
     "fast configuration", "[N]"],
    ["Pricing", "Recurring and one-time charges, discounts, fees, taxes", "[N]"],
    ["Business rules", "Eligibility by market, distribution channel and customer category; dependencies and "
     "constraints", "[N]"],
    ["Catalog federation", "Brings partner catalogs into one product management interface", "[N]"],
    ["Analytics", "Product consumption patterns used to target new offers", "[N]"],
    ["Open API layer", "Exposes the catalog to CRM, CPQ, ordering and billing through TM Forum APIs", "[N][S]"],
], widths=[4.6, 9.8, 2.2])

image("02_netcracker_catalog_feeds_sales_and_orders.png",
      "Figure 2. The Product Offering Catalog (commercial) and the Technical Product Catalog feed the sales "
      "side (CPQ) and Order Management. Source: Netcracker CPQ whitepaper (2021), Figure 2.")

h2("2.2 Core data model [S][N]")
para("TM Forum certified Netcracker's product data model for these entities. The example column uses one "
     "illustrative enterprise product that runs through the whole document [E].")
table([
    ["Entity", "Plain meaning", "Example [E]"],
    ["Product Specification", "What the product technically is, with its characteristics",
     "Dedicated Internet Access, bandwidth = 100 Mbps"],
    ["Product Offering", "What is sold: a single offer or a bundle", "\"Business Internet 100\" bundle"],
    ["Product Offering Price", "The prices attached to an offering", "Monthly fee, installation fee per site"],
    ["Product Catalog / Category", "A grouping of offerings", "Enterprise connectivity catalog"],
    ["Product", "What a customer actually owns after the sale", "Company X's Business Internet 100 at 3 sites"],
], widths=[4.2, 6.6, 5.8])

# ------------------------------------------------------------------ 3
h1("3. Using the system: from idea to launch-ready")
para("Each stage lists its status, owner, inputs and exit criteria, followed by what the user does in the "
     "system. The running example is \"Business Internet 100\": 100 Mbps dedicated internet, a block of "
     "static IP addresses, a managed router, and an availability SLA, sold on 12, 24 or 36-month contracts [E].")

stage(
    "3.1 Stage 1: Ideation and business case",
    "In Study [S]",
    "Capture a product idea, check its business value, and decide whether it is worth designing.",
    "Product Manager (owner), Marketing, Finance (business case)",
    "Market need, competitor offers, consumption analytics, target segment",
    [
        "Product Manager opens Product Lifecycle Management and registers a new concept: "
        "\"Business Internet 100 for SMEs\".",
        "Attaches the business case: target segment, expected number of sites, revenue goal.",
        "Reviews consumption analytics to size the opportunity (e.g. customers outgrowing 50 Mbps).",
        "Searches the catalog for existing specifications and offers that can be reused.",
        "Submits the concept for the conception approval gate.",
    ],
    "Concept approved, so the status moves to In Design; or the concept is dropped.",
    [
        "Netcracker PLM starts at \"marketing concept development\" [N].",
        "Netcracker analyzes \"product consumption patterns to enable better product targeting\" [N].",
        "TMF620: first status is In Study; \"conception approved\" moves it to In Design [S].",
    ],
    ["Business-case screens and who signs the conception approval."],
)

stage(
    "3.2 Stage 2: Technical product definition",
    "In Design [S]",
    "Define what the product is technically, so it can be delivered consistently.",
    "Product Specialist / Network & IT (owner), Product Manager (consulted)",
    "Approved concept, network capabilities, existing specifications",
    [
        "Product Specialist opens the Technical Product Catalog.",
        "Selects the \"Dedicated Internet Access\" specification (or a template) and sets its "
        "characteristics: bandwidth 100 Mbps, access type fiber, static IP block size.",
        "Adds the related specifications: managed router, IP address block, SLA monitoring.",
        "Defines decomposition rules: which technical tasks an order creates (site survey, last-mile "
        "installation, router configuration, IP allocation).",
        "Saves; the specifications are now available for commercial offers.",
    ],
    "Product specifications complete and linked to their delivery definitions.",
    [
        "Technical Product Catalog holds \"product/service specifications, provisioning templates, "
        "decomposition rules\" [N] (Figure 2).",
        "Product Specification is a TM Forum-certified entity in Netcracker's data model [S][N].",
        "Configuration \"using out-of-the-box product components\" [N].",
    ],
)

stage(
    "3.3 Stage 3: Commercial offer and bundling",
    "In Design [S]",
    "Package the technical product into something sales can sell.",
    "Product Manager (owner)",
    "Specifications from Stage 2, concept from Stage 1",
    [
        "Product Manager opens the Product Offering Catalog and creates a new offering from a template.",
        "Links it to the Dedicated Internet Access specification.",
        "Builds the bundle: internet 100 Mbps + static IP block + managed router + SLA.",
        "Adds optional add-ons (e.g. extra IP block, backup link) as separate offerings that can be "
        "attached.",
        "Places the offering in the Enterprise Connectivity catalog.",
    ],
    "Offering structure complete and linked to its specifications.",
    [
        "Product Offering Catalog holds \"offering details, business rules, prices and discounts\" [N].",
        "\"Centralizing the management of products and bundles\" [N].",
    ],
)

stage(
    "3.4 Stage 4: Pricing",
    "In Design [S]",
    "Attach all prices and price-related terms to the offering.",
    "Product Manager (owner), Finance (approver)",
    "Offering from Stage 3, pricing strategy, tax rules",
    [
        "Product Manager opens the offering's pricing section.",
        "Adds a recurring monthly fee and a one-time installation fee per site.",
        "Sets prices by contract term (12, 24, 36 months) and by customer category.",
        "Adds discounts (e.g. multi-site discount) and applies taxes.",
        "Sets the price validity period.",
    ],
    "Complete price set attached to the offering.",
    [
        "Catalog covers \"pricing, discounts, fees, loyalty programs, taxes\" [N].",
        "Pricing \"based on market, customer category, distribution channel\" (CPQ whitepaper) [N].",
        "Product Offering Price is a TM Forum-certified entity [S][N].",
    ],
    ["Price-change approval rules and contract-term pricing screens."],
)

stage(
    "3.5 Stage 5: Business rules and eligibility",
    "In Design [S]",
    "Control who can buy the offer, where, through which channel, and with what.",
    "Product Manager (owner), Product Specialist (consulted)",
    "Offering and prices; coverage; channel strategy",
    [
        "Product Manager opens the offering's rules.",
        "Sets eligibility: market = fiber-covered areas, channel = enterprise account managers, "
        "customer category = SME and corporate.",
        "Adds dependencies: an extra IP block requires the base offering.",
        "Adds constraints: the backup link add-on is only sold with a 24 or 36-month contract.",
    ],
    "Rules defined, so CRM will only offer the product to eligible customers.",
    [
        "Validation against \"business rules, such as eligibility rules, configured in the Product catalog\" "
        "to check availability \"for the market, distribution channel and customer category\" [N].",
        "Rules cover \"product specifications, dependencies, constraints and business rules\" [N].",
    ],
)

stage(
    "3.6 Stage 6: Approval gates and roles",
    "In Design to In Test [S]",
    "Formally approve the design before testing starts.",
    "Product Manager (submits); Marketing head, Finance, Legal / Regulatory, IT/Network (approvers) [E]",
    "Complete design from Stages 2 to 5",
    [
        "Product Manager submits the offering for design approval.",
        "Approvers review: commercial fit (Marketing), prices and margin (Finance), terms and regulation "
        "(Legal), technical feasibility (IT/Network).",
        "Each approver approves or rejects with comments, which are kept with the product.",
        "When all approve, the status changes to In Test; if rejected, it stays In Design for rework.",
    ],
    "Design approved and status In Test.",
    [
        "TMF620 defines the gates: \"conception approved\" to In Design and \"design approved\" to In Test [S].",
        "Netcracker's tools \"streamline approval workflows\" and show approval statuses (Figure 3) [N].",
        "PLM provides \"automating end-to-end lifecycle management\" [N].",
    ],
    ["Approval workflow configuration, role permissions, and sequential vs parallel approval."],
)

stage(
    "3.7 Stage 7: Versioning",
    "Any status [S]",
    "Keep a controlled history of every change to the product.",
    "Product Manager, Catalog Administrator [E]",
    "An existing product or offering that needs to change",
    [
        "Instead of editing an approved product directly, the user creates a new version.",
        "The new version goes back through design, approval and test.",
        "Previous versions stay visible for comparison and audit.",
    ],
    "New version approved and tested; old version kept in history.",
    [
        "PLM \"fully supports product versioning, upgrades\" [N].",
        "TMF620 catalog items carry a version and a lifecycle status [S].",
    ],
    ["Version comparison screens and whether two versions can be in design at once."],
)

stage(
    "3.8 Stage 8: Testing and readiness",
    "In Test to Active (or Rejected) [S]",
    "Prove the product works end to end before any customer can buy it.",
    "Test team (owner), IT, Product Manager",
    "Approved design",
    [
        "The approved product is deployed to a test environment.",
        "Test team runs test quotes and orders to check pricing, rules and decomposition.",
        "If tests pass, the product moves to Active, ready to be published to CRM.",
        "If tests fail, the product is Rejected (final) or reworked as a new version.",
        "Optional test launch to a limited group of customers.",
    ],
    "Status Active: validated and ready to hand over to CRM.",
    [
        "PLM covers \"product testing, test launch and subsequent mass production\" [N].",
        "TMF620: test OK gives Active; test KO gives Rejected, a final status [S].",
    ],
)

h2("3.9 The journey on one page")
table([
    ["#", "Stage", "Status", "Netcracker module", "Owner", "Output"],
    ["1", "Ideation", "In Study", "PLM + analytics", "Product Manager", "Approved concept"],
    ["2", "Technical definition", "In Design", "Technical Product Catalog", "Product Specialist", "Specifications"],
    ["3", "Commercial offer", "In Design", "Product Offering Catalog", "Product Manager", "Offering / bundle"],
    ["4", "Pricing", "In Design", "Catalog pricing", "PM + Finance", "Price set"],
    ["5", "Rules & eligibility", "In Design", "Business rules", "Product Manager", "Eligibility rules"],
    ["6", "Approval", "In Design to In Test", "PLM", "Approvers", "Approved design"],
    ["7", "Versioning", "Any", "PLM", "PM / Admin", "Version history"],
    ["8", "Testing", "In Test to Active", "PLM", "Test team / IT", "Launch-ready product"],
], widths=[0.7, 3.2, 2.8, 3.6, 2.9, 3.4])

# ------------------------------------------------------------------ 4
h1("4. Integration with CRM: how Netcracker hands the product over")

h2("4.1 The principle: CRM consumes, it does not re-create")
para("In Netcracker's model the catalog is the single source of truth. CRM does not create its own copy of "
     "a product; it reads the approved definition, including prices and rules, from the catalog. The "
     "Product Catalog is described as \"supporting sales, ordering and billing processes via "
     "industry-standard integration\" [N], and Netcracker's lead-to-order process is \"automated with a "
     "catalog-driven approach\" [N].")

h2("4.2 What gets published")
table([
    ["Item", "Why CRM needs it", "Source"],
    ["Offering and bundle structure", "So the salesperson can select the product and its add-ons", "[N][S]"],
    ["Prices, discounts, contract terms", "So the quote price is always the approved price", "[N][S]"],
    ["Eligibility rules (market, channel, customer category)", "So CRM only shows products the customer can buy", "[N]"],
    ["Dependencies and constraints", "So invalid combinations cannot be quoted", "[N]"],
    ["Validity dates (start and end of sale)", "So the product appears and disappears on time", "[S]"],
    ["Lifecycle status and version", "So CRM knows which version is sellable", "[S]"],
], widths=[5.4, 8.8, 2.4])

h2("4.3 When it gets published")
bullets([
    "Only products in Active status are published; anything In Study, In Design or In Test is invisible "
    "to sales [S].",
    "The start of marketing (Active to Launched) can be tied to the offering's validity start date, so a "
    "product can be prepared in advance and go on sale on a planned date [S].",
    "Publication can be limited to chosen channels and customer categories through the eligibility rules "
    "defined in Stage 5 [N].",
])

h2("4.4 How it is published: the TM Forum Open APIs")
table([
    ["API", "What it does in the hand-over", "Netcracker certified"],
    ["TMF620 Product Catalog Management", "CRM reads catalogs, offerings, prices and specifications", "Yes [N]"],
    ["TMF620 event notifications", "CRM subscribes and is notified when an offering is created, changed or "
     "changes status, so it stays in sync without manual updates", "Part of TMF620 [S]"],
    ["TMF679 Product Offering Qualification", "At quote time CRM asks \"can this customer get this offer "
     "here?\" and the catalog rules answer", "Standard [S]"],
    ["TMF648 Quote Management", "Quotes are built on catalog offerings and prices", "Yes [N]"],
    ["TMF622 Product Ordering", "The order references the offering ID and version, so delivery and "
     "billing use exactly what was approved", "Yes [N]"],
], widths=[4.6, 8.4, 3.6])

h2("4.5 What the salesperson sees in CRM")
para("When a salesperson creates a quote, the catalog drives the screen. The fields Customer Category, "
     "Distribution Channel and Market on Netcracker's Create Quote screen are the same eligibility "
     "dimensions defined in Stage 5, and the offer is validated against the catalog rules before it can be "
     "submitted [N].")
image("03_netcracker_real_screen_create_quote.png",
      "Figure 3. Real Netcracker screen (Create Quote). Customer Category, Distribution Channel and Market "
      "come from the catalog's eligibility rules. Source: Netcracker CPQ whitepaper (2021), Figure 3.",
      width_cm=13)

h2("4.6 When a product changes after launch")
numbered([
    "The Product Manager creates a new version in the catalog (Stage 7) and takes it through approval and test.",
    "When the new version becomes Active, a change event is sent to CRM [S].",
    "New quotes and orders use the new version; existing customers keep the version they bought, because "
    "each order recorded the offering version [S].",
    "When the old version reaches its end-of-sale date, it is Retired and disappears from CRM sales screens [S].",
])

h2("4.7 Example: connecting the catalog to Microsoft Dynamics 365")
para("The same pattern applies to a CRM from another vendor. Dynamics 365 Sales has its own product "
     "entities [M]. There is no ready-made Netcracker connector for Dynamics, so the link would be built "
     "through an integration layer that calls the TMF620 API on one side and the Dynamics (Dataverse) Web "
     "API on the other [E].")
table([
    ["Catalog item (TMF620)", "Dynamics 365 Sales entity [M]"],
    ["Product Offering", "Product"],
    ["Bundle offering", "Product bundle"],
    ["Product Catalog / Category", "Product family"],
    ["Product Offering Price", "Price list item (within a price list)"],
    ["Discounts", "Discount list"],
    ["Related offers / add-ons", "Product relationships (accessory, cross-sell, up-sell)"],
    ["Lifecycle status: Active / Launched / Retired", "Product status: Active / Retired"],
], widths=[7.6, 9.0])
para("Dynamics product statuses are Draft, Active, Under Revision and Retired [M]. In this pattern the "
     "statuses before Active (study, design, test and approval) stay in the catalog, and Dynamics receives "
     "only products that are ready to sell.", italic=True)

# ------------------------------------------------------------------ 5
h1("5. Evidence from real operators")
table([
    ["Operator", "What was deployed", "Published result", "Source type"],
    ["Zain KSA (Saudi Arabia)", "Full BSS/OSS including Product Catalog",
     "50% decrease in product development time; delivered in under 3 years",
     "Netcracker case study with operator quotes"],
    ["T-Mobile Netherlands", "Product Catalog + Product Lifecycle Management",
     "Manages products \"from development to production to retirement\"; fewer incorrect offers",
     "Netcracker press release with operator quote"],
    ["e& UAE", "Full BSS including product catalog management", "\"Groundwork for catalog-driven offers\"",
     "Netcracker press release with operator quote"],
], widths=[3.4, 4.4, 5.0, 3.8])
image("04_zain_ksa_case_study_page1.png",
      "Figure 4. Zain KSA case study: 50% decrease in product development time. Source: Netcracker.", width_cm=12)

# ------------------------------------------------------------------ 6
h1("6. Points to explore further")
bullets([
    "How ideas and business cases are captured and scored before In Study.",
    "How the approval workflow is configured: approvers per stage, sequential or parallel, delegation.",
    "Which roles and permissions exist for catalog users (view, create, edit, approve, publish).",
    "Whether custom lifecycle statuses can be added (e.g. a regulatory approval step).",
    "How versions are compared, and whether two versions can be in design at the same time.",
    "How catalog events are delivered to an external CRM, and how failed synchronizations are handled.",
    "How contract-term and multi-site pricing for enterprise offers are modelled.",
])

# ------------------------------------------------------------------ 7
h1("7. References")
refs = [
    "Netcracker, Product Management (Product Catalog + PLM): "
    "https://www.netcracker.com/portfolio/products/digital-bss/product-management",
    "Netcracker, Commerce Management: "
    "https://www.netcracker.com/portfolio/products/digital-commerce-monetization/commerce-management",
    "Netcracker, Configure-Price-Quote whitepaper (2021), Figures 2 and 3: "
    "https://www.netcracker.com/assets/portfolio/solutions/intelligent-sales-automation/netcracker_cpq-whitepaper.pdf",
    "Netcracker, Zain KSA case study: https://netcracker.com/documents/pdf/netcracker_zain_case-study_260918.pdf",
    "Netcracker, T-Mobile Netherlands press release: "
    "https://www.netcracker.com/news/press-releases/t-mobile-netherlands-optimizes-end-to-end-product-management-with-netcracker-bss/",
    "Netcracker, e& UAE partnership (Oct 2023): "
    "https://netcracker.com/news/press-releases/etisalat-by-e-and-partners-with-netcracker-on-the-largest-full-stack-bss-transformation-project-in-the-middle-east",
    "TM Forum, Netcracker SID certification report (Figure 1): "
    "https://www.tmforum.org/wp-content/uploads/2015/03/Netcrakcer-Conformance-Cert-NetCracker-TOMS-Suite-Release-8.2-SID-Report.pdf",
    "TM Forum, Open API conformance list (Netcracker certifications): "
    "https://www.tmforum.org/conformance-certification/open-api-conformance/",
    "TM Forum, TMF620 Product Catalog Management API: "
    "https://www.tmforum.org/open-digital-architecture/open-apis/product-catalog-management-api-TMF620/v5.0",
    "TM Forum, TMF679 Product Offering Qualification API: "
    "https://www.tmforum.org/open-digital-architecture/open-apis/product-offering-qualification-api-TMF679",
    "Microsoft Learn, Set up a product catalog in Dynamics 365 Sales: "
    "https://learn.microsoft.com/en-us/dynamics365/sales/set-up-product-catalog-walkthrough",
]
numbered(refs)

doc.save(OUT)
print("saved", OUT.name)
