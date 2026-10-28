"""Builds the PLC study document (Amdocs CatalogONE, catalog-only PLC that publishes to consuming apps / CRM).

Run:  python3 build_document_catalogone.py
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).parent
IMG = HERE / "images" / "catalogone"
OUT = HERE / "PLC_CatalogONE_Study.docx"

NAVY = RGBColor(0x0C, 0x29, 0x4B)
GREY = RGBColor(0x5F, 0x5A, 0x6B)
HDR = "0C294B"

doc = Document()
for section in doc.sections:
    section.left_margin = section.right_margin = Cm(2.2)
    section.top_margin = section.bottom_margin = Cm(2.0)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
    st = doc.styles[name]
    st.font.name = "Calibri"
    st.font.size = Pt(size)
    st.font.color.rgb = NAVY


def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def h1(t):
    doc.add_heading(t, level=1)


def h2(t):
    doc.add_heading(t, level=2)


def h3(t):
    doc.add_heading(t, level=3)


def para(t, italic=False, color=None, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
    r = p.add_run(t)
    r.italic = italic
    if color:
        r.font.color.rgb = color
    return p


def bullets(items, style="List Bullet"):
    for it in items:
        p = doc.add_paragraph(style=style)
        if isinstance(it, tuple):
            r = p.add_run(it[0])
            r.bold = True
            p.add_run(it[1])
        else:
            p.add_run(it)


def numbered(items):
    for i, it in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(f"{i}.  ")
        r.bold = True
        r.font.color.rgb = NAVY
        p.add_run(it)


def table(rows, widths=None, header_fill=HDR):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            cell.text = ""
            run = cell.paragraphs[0].add_run(val)
            run.font.size = Pt(9.5)
            if r == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                shade(cell, header_fill)
            elif c == 0:
                run.bold = True
            if widths:
                cell.width = Cm(widths[c])
            if r < len(rows) - 1 and len(rows) <= 12:
                cell.paragraphs[0].paragraph_format.keep_with_next = True
    if widths:
        t.autofit = False
        for c, w in enumerate(widths):
            t.columns[c].width = Cm(w)
    doc.add_paragraph()
    return t


def note(title, text):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    cell = t.cell(0, 0)
    shade(cell, "E8F0F8")
    p = cell.paragraphs[0]
    r = p.add_run(title + "  ")
    r.bold = True
    r.font.color.rgb = NAVY
    p.add_run(text).font.size = Pt(10)
    doc.add_paragraph()


def image(file, caption, width_cm=16):
    doc.add_picture(str(IMG / file), width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(caption)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = GREY


def stage(title, status, purpose, who, inputs, steps, outputs, evidence, gaps=None):
    h2(title)
    table([
        ["Item", "Detail"],
        ["Lifecycle status", status],
        ["Purpose", purpose],
        ["Who (proposed roles) [P]", who],
        ["Inputs", inputs],
        ["Exit criteria / output", outputs],
    ], widths=[4.2, 12.4])
    h3("What you do in the system (walkthrough)")
    numbered(steps)
    h3("Evidence")
    bullets(evidence)
    if gaps:
        h3("Not publicly documented (confirm in a vendor demo)")
        bullets(gaps)


# ------------------------------------------------------------------ cover
t = doc.add_paragraph()
r = t.add_run("PLC STUDY  |  WE DATA")
r.bold = True
r.font.color.rgb = NAVY
r.font.size = Pt(11)
title = doc.add_paragraph()
r = title.add_run("Product Lifecycle Management in Amdocs CatalogONE:\nfrom Idea to CRM")
r.bold = True
r.font.size = Pt(22)
para("System studied: Amdocs CatalogONE, a standalone telecom product catalog and "
     "Product Lifecycle Management (PLM) platform. It is BSS-agnostic: it owns the "
     "catalog, not the CRM. Amdocs documents that CatalogONE publishes to consuming "
     "applications and to \"the customer's CRM\". The Amdocs–Microsoft Customer "
     "Engagement Platform (CEP) is the product that puts an Amdocs Catalog next to "
     "Dynamics 365. Amdocs does not publish a CatalogONE-branded connector named "
     "Dynamics 365.", color=GREY)
para("Prepared by: [Your name]    |    Version: Draft 4 for review    |    October 2026",
     color=GREY)

note("How to read this document.",
     "Every statement is tagged: [A] = Amdocs public datasheets, whitepapers or "
     "the Amdocs–Microsoft Customer Engagement Platform; [P] = Appledore Research "
     "solution profile of CatalogONE (2020); [T] = TM Forum case study or Open API "
     "conformance list; [S] = TM Forum TMF620 and related Open APIs; [M] = Microsoft "
     "documentation; [I] = IDC customer case study on PLDT Unified Case Management "
     "(hosted by Amdocs, June 2026); [E] = illustrative example. CatalogONE user guides are not "
     "public, so screen names in the walkthroughs are illustrative and will differ "
     "in a real installation.")

# ------------------------------------------------------------------ 0
h1("Key takeaways")
bullets([
    ("Catalog only, not a CRM. ", "CatalogONE is a product catalog and PLM. It is "
     "not a CRM. It publishes approved catalog data to consuming applications; "
     "Amdocs names \"the customer's CRM\" as one of those consumers [A]."),
    ("One source of truth. ", "Commercial, charging and network-service definitions "
     "live in one catalog. Downstream systems consume it instead of keeping their "
     "own copies [A][P]."),
    ("A governed journey. ", "PLM covers ideation through creation, implementation "
     "and end-of-life, with workflow, approvals, simulations, parallel work and "
     "versioned change [A]. Statuses follow the TM Forum catalog model [S]."),
    ("Business users configure, not IT. ", "A role-based UI, templates and building "
     "blocks let product managers create offerings, prices, channels and segments "
     "with no code and no deployment overhead [A]."),
    ("Publish, do not re-key. ", "CatalogONE Centralized authors once and publishes "
     "catalog information to consuming applications. Distribution uses a Kafka bus "
     "and TMF APIs, in minutes rather than weeks [A][P]. Dynamics 365 is not named "
     "on the CatalogONE datasheet; it is named on the Amdocs–Microsoft CEP [A]."),
    ("Proven results. ", "A European quad-play operator moved new-offer time from "
     "up to nine months to 1–2 weeks. Sprint / T-Mobile expected a 50% time-to-market "
     "cut with CatalogONE. KT launched CatalogONE to accelerate 5G offers [P][T][A]."),
])

# ------------------------------------------------------------------ 1
doc.add_page_break()
h1("1. PLC and the TM Forum standard in brief")
para("A Product Lifecycle Management (PLC) system is where an operator turns an idea "
     "into a sellable product, controls every change to it, and publishes it to the "
     "systems that sell, deliver and bill it.")
para("TM Forum defines this lifecycle in the Product Catalog Management API (TMF620). "
     "Each catalog item carries one of these statuses [S]:")
table([
    ["Status", "Meaning", "Gate to enter it"],
    ["In Study", "The idea is being studied", "Idea registered"],
    ["In Design", "The product is being designed", "Conception approved"],
    ["In Test", "The product is being tested", "Design approved"],
    ["Rejected", "Test failed (final status)", "Test KO"],
    ["Active", "Ready, but not yet sold", "Test approved"],
    ["Launched", "Available to customers through sales channels", "Start of marketing"],
    ["Retired / Obsolete", "No new sales / no customers left", "End of marketing"],
], widths=[3.2, 8.2, 5.2])
para("CatalogONE is built on the TM Forum data model and exposes TMF Open APIs. "
     "Amdocs is listed on the TM Forum Open API conformance register for TMF620 "
     "(Product Catalog Management) and related commerce APIs. CatalogONE PLM is "
     "described as covering ideation through end-of-life, so the stages below follow "
     "the industry model rather than a vendor-only one [A][T][S]. This document covers "
     "everything up to Launched: the stages before CRM, how CatalogONE publishes "
     "to consuming systems, and how that relates to Dynamics 365 in the "
     "Amdocs–Microsoft platform.", italic=True)

# ------------------------------------------------------------------ 2
h1("2. CatalogONE at a glance")
para("CatalogONE is Amdocs' cloud-native catalog. It is designed as a stand-alone "
     "application that can sit on top of existing catalogs and BSS systems, or act "
     "as the master catalog. It does not replace the CRM [A][P]. Business users "
     "configure products, prices, channels and segments with no code; PLM includes "
     "workflow and approval management [A].")

h2("2.1 What it is, and what it is not")
table([
    ["It is", "It is not"],
    ["A telecom product catalog and PLM", "A CRM"],
    ["BSS-agnostic and API-first [A]", "Tied to replacing Dynamics"],
    ["The authoring and approval system for offers", "The quote, order or billing screen"],
    ["Optional Charging catalog for rated usage [A]", "Required for unrated enterprise products"],
], widths=[8.3, 8.3])
para("For an operator that already sells unrated enterprise products in Microsoft "
     "Dynamics 365, the relevant CatalogONE parts are Commerce (offers and bundles), "
     "Centralized (master catalog and publish), Product Lifecycle Management, and "
     "the TMF API / event layer. The Charging catalog is an optional module for "
     "operators that rate usage; it is not required for this pattern [A][E].")

h2("2.2 Modules")
para("CatalogONE is modular. Commerce, Charging, Centralized, Product Lifecycle "
     "Management and iCatalog can be adopted one priority at a time [A].")
table([
    ["Module", "What it does", "Source"],
    ["Product Lifecycle Management",
     "Automation and governance from ideation through creation, implementation "
     "and end-of-life; simulations, parallel work, workflow and approvals", "[A]"],
    ["CatalogONE Commerce",
     "Standalone commercial catalog: offerings, bundles, promotions, channels, "
     "segments; role-based UI and templates; can sit as a layer on existing catalogs", "[A]"],
    ["CatalogONE Centralized",
     "Master catalog: unified data model, collaborative authoring, testing, "
     "A/B testing, auto-release; publishes to consuming applications", "[A]"],
    ["Service / resource catalog",
     "Technical catalog: network and IT capabilities packaged for business users; "
     "maps down to service and resource definitions", "[P]"],
    ["Templates and building blocks",
     "Pre-defined product templates (classic telco, infotainment, licensed products, "
     "hard goods) and reusable components", "[A][P]"],
    ["Pricing and policy",
     "Prices, promotions, entitlements, channel and segment rules; Charging catalog "
     "adds rated usage logic where needed", "[A]"],
    ["Catalog federation",
     "Bring partner and third-party catalogs into one interface via TMF Open APIs", "[A]"],
    ["iCatalog (intelligence)",
     "KPI dashboards, predictive analytics, recommendations, cannibalisation alerts", "[A]"],
    ["Distribution / Open API layer",
     "Kafka bus plus TMF APIs to synchronise approved catalog data to CRM, CPQ, "
     "ordering and billing in minutes", "[P][A]"],
], widths=[4.6, 9.8, 2.2])

h2("2.3 Commercial catalog, technical catalog, master catalog")
para("CatalogONE keeps three views of the same product, then publishes them [A][P]:")
bullets([
    ("Commercial catalog: ", "what is sold — offerings, bundles, prices, promotions, "
     "channels and segments. Owned by the business user [A]."),
    ("Service / resource catalog: ", "how it is delivered — network and IT capabilities "
     "exposed so marketing can reuse them without changing the network side [P]."),
    ("Centralized / master catalog: ", "one data model. Author once, then publish to "
     "every consuming application (CRM, billing, charging, order management) [A]."),
])
image("06_centralized_publish.png",
      "Figure 1. CatalogONE Centralized: collaborative authoring, testing automation, "
      "A/B testing and auto-release, then publish catalog information to consuming "
      "applications. Source: Amdocs 6 Steps ebook, p.15 [A].", width_cm=14)

h2("2.4 Core data model [S][A][P]")
para("CatalogONE's APIs are based on the TM Forum data model [A]. The example column "
     "uses one illustrative enterprise product that runs through the whole document [E].")
table([
    ["Entity", "Plain meaning", "Example [E]"],
    ["Product Specification", "What the product technically is, with its characteristics",
     "Dedicated Internet Access, bandwidth = 100 Mbps"],
    ["Product Offering", "What is sold: a single offer or a bundle",
     "\"Business Internet 100\" bundle"],
    ["Product Offering Price", "The prices attached to an offering",
     "Monthly fee, installation fee per site"],
    ["Product Catalog / Category", "A grouping of offerings",
     "Enterprise connectivity catalog"],
    ["Channel / segment / entitlement", "Where and to whom the offer is sold",
     "Enterprise account managers; SME and corporate"],
    ["Product", "What a customer actually owns after the sale",
     "Company X's Business Internet 100 at 3 sites"],
], widths=[4.4, 6.4, 5.8])

# ------------------------------------------------------------------ 3
h1("3. Using the system: from idea to launch-ready")
para("Each stage lists its status, owner, inputs and exit criteria, followed by what "
     "the user does in the system. The running example is \"Business Internet 100\": "
     "100 Mbps dedicated internet, a block of static IP addresses, a managed router, "
     "and an availability SLA, sold on 12, 24 or 36-month contracts [E].")
para("Amdocs describes a role-based UI with guided flows: marketing / business users "
     "own commercial changes (prices, promotions, bundles); IT, network and product "
     "managers work in the same design portal with a different personalised "
     "experience [A]. Screen names below are illustrative [E].")

image("03_six_steps_text.png",
      "Figure 2. CatalogONE's six catalog capabilities. Step 5 is Product Lifecycle "
      "Management across the organisation. Source: Amdocs 6 Steps ebook, p.7 [A].",
      width_cm=10)

stage(
    "3.1 Stage 1: Ideation and business case",
    "In Study [S]",
    "Capture a product idea, check its business value, and decide whether it is worth designing.",
    "Product Manager (owner), Marketing, Finance (business case)",
    "Market need, competitor offers, iCatalog analytics, target segment",
    [
        "Product Manager opens CatalogONE PLM and registers a new concept: "
        "\"Business Internet 100 for SMEs\".",
        "Attaches the business case: target segment, expected number of sites, revenue goal.",
        "Reviews iCatalog dashboards: which current offers are under-performing, which "
        "segments are growing, whether a 100 Mbps offer would cannibalise 50 Mbps [A].",
        "Searches the commercial and technical catalogs for specifications and offers "
        "that can be reused or federated.",
        "Submits the concept for the conception approval gate.",
    ],
    "Concept approved, so the status moves to In Design; or the concept is dropped.",
    [
        "PLM starts at ideation [A].",
        "iCatalog \"reduces the new offer ideation cycle time, as well as the risk of "
        "rolling out sub-optimal offers\" [A].",
        "TMF620: first status is In Study; \"conception approved\" moves it to In Design [S].",
    ],
    ["Business-case scoring screens and who signs the conception approval."],
)

image("07_icatalog_dashboard.png",
      "Figure 3. iCatalog: dashboards and analytics used during ideation and after "
      "launch to optimise offers. Source: Amdocs 6 Steps ebook, p.16 [A].",
      width_cm=12)

stage(
    "3.2 Stage 2: Technical product definition",
    "In Design [S]",
    "Define what the product is technically, so it can be delivered consistently.",
    "Product Specialist / Network & IT (owner), Product Manager (consulted)",
    "Approved concept, network capabilities, existing specifications",
    [
        "Network / IT opens the service / resource catalog in the same design portal, "
        "on the technical role view [A][P].",
        "Selects the \"Dedicated Internet Access\" specification (or a template) and "
        "sets characteristics: bandwidth 100 Mbps, access type fibre, static IP block size.",
        "Adds related specifications: managed router, IP address block, SLA monitoring.",
        "Exposes those capabilities to the commercial catalog so the Product Manager "
        "can bundle them without changing the network side [P].",
        "Saves; the specifications are now available for commercial offers.",
    ],
    "Product specifications complete and linked to their delivery definitions.",
    [
        "Role-based UI: \"Network can create new services; IT new products; and the "
        "business can define new offers and bundles\" [P].",
        "Service / resource catalog packages network and IT capabilities for business "
        "owners [P].",
        "Product Specification is a TM Forum catalog entity [S].",
    ],
)

stage(
    "3.3 Stage 3: Commercial offer and bundling",
    "In Design [S]",
    "Package the technical product into something sales can sell.",
    "Product Manager (owner)",
    "Specifications from Stage 2, concept from Stage 1",
    [
        "Product Manager opens CatalogONE Commerce and starts the offer wizard from "
        "a template [A].",
        "Links the offer to the Dedicated Internet Access specification.",
        "Builds the bundle: internet 100 Mbps + static IP block + managed router + SLA.",
        "Adds optional add-ons (extra IP block, backup link) as separate offerings.",
        "Places the offering in the Enterprise Connectivity catalog and selects "
        "channels and segments [A].",
    ],
    "Offering structure complete and linked to its specifications.",
    [
        "Commerce is a standalone application with templates and building blocks for "
        "promotions and bundles [A].",
        "Business users \"configure new products and services, pricing, channels, "
        "segments, and more, without any code\" [A].",
        "Configure once, then launch to all channels (self-service, mobile, B2B portal) [A].",
    ],
)

stage(
    "3.4 Stage 4: Pricing",
    "In Design [S]",
    "Attach all prices and price-related terms to the offering.",
    "Product Manager (owner), Finance (approver)",
    "Offering from Stage 3, pricing strategy, tax rules",
    [
        "Product Manager opens the offering's pricing section in the wizard.",
        "Adds a recurring monthly fee and a one-time installation fee per site.",
        "Sets prices by contract term (12, 24, 36 months) and by customer segment.",
        "Adds discounts (for example a multi-site discount) and validity dates.",
        "If the operator later needs rated usage, the Charging catalog can attach "
        "rate tables; it is not required for unrated enterprise products [A][E].",
    ],
    "Complete price set attached to the offering.",
    [
        "Catalog covers \"new pricing definition, promotions, rules, entitlements\" [A].",
        "Charging catalog: implement, extend, validate and approve pricing logic "
        "and policies with zero coding — optional module [A].",
        "Product Offering Price is a TM Forum catalog entity [S].",
    ],
    ["Price-change approval rules and contract-term pricing screens."],
)

stage(
    "3.5 Stage 5: Business rules and eligibility",
    "In Design [S]",
    "Control who can buy the offer, where, through which channel, and with what.",
    "Product Manager (owner), Product Specialist (consulted)",
    "Offering and prices; coverage; channel strategy",
    [
        "Product Manager opens the offering's rules in the wizard.",
        "Sets eligibility: market = fibre-covered areas, channel = enterprise account "
        "managers, segment = SME and corporate.",
        "Adds dependencies: an extra IP block requires the base offering.",
        "Adds constraints: the backup-link add-on is only sold with a 24 or 36-month contract.",
        "Sets KPIs for iCatalog to watch after launch [A].",
    ],
    "Rules defined, so the sales channel / CPQ will only offer the product to eligible customers.",
    [
        "Wizard includes \"channel selection, segment selection and setting KPIs\" "
        "before launch across those channels and segments [A].",
        "Marketers can \"view dependencies in order to understand offering constraints\" [A].",
        "TMF679 Product Offering Qualification is the standard \"can this customer "
        "get this offer here?\" API; Sprint used it with Amdocs DigitalONE [T][S].",
    ],
)

stage(
    "3.6 Stage 6: Approval gates and roles",
    "In Design to In Test [S]",
    "Formally approve the design before testing starts.",
    "Product Manager (submits); Marketing head, Finance, Legal / Regulatory, "
    "IT/Network (approvers) [E]",
    "Complete design from Stages 2 to 5",
    [
        "Product Manager submits the offering for design approval in PLM.",
        "Approvers review: commercial fit (Marketing), prices and margin (Finance), "
        "terms and regulation (Legal), technical feasibility (IT/Network).",
        "Each approver approves or rejects with comments, kept with the product.",
        "When all approve, the status changes to In Test; if rejected, it stays "
        "In Design for rework.",
    ],
    "Design approved and status In Test.",
    [
        "PLM supports \"workflow and approval management\" and \"effective collaboration "
        "both within and outside the organization\" [A].",
        "KT: CatalogONE provides \"approvals and notification management capabilities\" "
        "so business and marketing users manage the offering lifecycle [P].",
        "TMF620 defines the gates: \"conception approved\" to In Design and "
        "\"design approved\" to In Test [S].",
    ],
    ["Approval workflow configuration, role permissions, sequential vs parallel."],
)

stage(
    "3.7 Stage 7: Versioning",
    "Any status [S]",
    "Keep a controlled history of every change to the product.",
    "Product Manager, Catalog Administrator [E]",
    "An existing product or offering that needs to change",
    [
        "Instead of editing a launched product directly, the user creates a new version.",
        "The new version goes back through design, approval and test (short iterative "
        "cycles; parallel work is supported) [A].",
        "Previous versions stay visible for comparison, audit and existing contracts.",
        "On approval, auto-release publishes the new version to consuming systems [A].",
    ],
    "New version approved and tested; old version kept in history.",
    [
        "PLM manages \"all changes that occur during the product lifecycle with short "
        "iterative cycles, working in agile mode\" [A].",
        "Governed workflows, approvals, validations and versioning \"allow CSPs to "
        "manage live offer changes safely\" [A].",
        "TMF620 catalog items carry a version and a lifecycle status [S].",
    ],
    ["Version comparison screens and whether two versions can be in design at once."],
)

stage(
    "3.8 Stage 8: Testing and readiness",
    "In Test to Active (or Rejected) [S]",
    "Prove the product works end to end before any customer can buy it.",
    "Test team (owner), IT, Product Manager",
    "Approved design",
    [
        "The approved product is released to a test / sandbox environment. CatalogONE "
        "distribution supports sandboxes, validation and testing [P].",
        "Test team runs test quotes and orders to check pricing, rules and delivery mapping.",
        "Centralized catalog supports testing automation and A/B testing [A].",
        "If tests pass, the product moves to Active, ready to be published to "
        "consuming applications (CRM, CPQ, ordering) [A].",
        "If tests fail, the product is Rejected (final) or reworked as a new version.",
        "Launch to production is described as \"at the touch of a button\" once "
        "approved [A].",
    ],
    "Status Active: validated and ready to publish to consuming applications.",
    [
        "Distribution includes \"sandboxes, validation and testing\" so a business-owner "
        "approval process can replace an IT-testing bottleneck [P].",
        "Centralized: testing automation, A/B testing, auto-release management [A].",
        "TMF620: test OK gives Active; test KO gives Rejected, a final status [S].",
    ],
)

h2("3.9 The journey on one page")
table([
    ["#", "Stage", "Status", "CatalogONE module", "Owner", "Output"],
    ["1", "Ideation", "In Study", "PLM + iCatalog", "Product Manager", "Approved concept"],
    ["2", "Technical definition", "In Design", "Service / resource catalog",
     "Product Specialist", "Specifications"],
    ["3", "Commercial offer", "In Design", "Commerce", "Product Manager", "Offering / bundle"],
    ["4", "Pricing", "In Design", "Commerce / Charging (optional)", "PM + Finance", "Price set"],
    ["5", "Rules & eligibility", "In Design", "Commerce wizard", "Product Manager", "Eligibility rules"],
    ["6", "Approval", "In Design to In Test", "PLM", "Approvers", "Approved design"],
    ["7", "Versioning", "Any", "PLM", "PM / Admin", "Version history"],
    ["8", "Testing", "In Test to Active", "PLM + Centralized", "Test team / IT", "Launch-ready product"],
], widths=[0.7, 3.2, 2.6, 4.0, 2.9, 3.2])

# ------------------------------------------------------------------ 4
doc.add_page_break()
h1("4. How the catalog reaches CRM — and what Amdocs actually documents")

note("Read this section first.",
     "Two facts are documented, and they are not the same sentence. "
     "(1) CatalogONE publishes catalog data to consuming applications and to "
     "\"the customer's CRM\" [A]. (2) The Amdocs–Microsoft Customer Engagement "
     "Platform puts an Amdocs unified Catalog next to Dynamics 365 Sales [A]. "
     "Amdocs does not publish a sentence that says \"CatalogONE sends data to "
     "Microsoft Dynamics 365 CRM\". The mapping table later in this section is "
     "an example [E], not a vendor connector spec.")

h2("4.1 What CatalogONE documents: publish to consuming systems, including CRM")
para("CatalogONE Centralized \"enables centralized authoring … publishing catalog "
     "information to multiple consuming applications\" [A]. The CatalogONE brochure "
     "describes \"integration to Amdocs CPQ and the customer's CRM solution\" [A]. "
     "It does not name Microsoft Dynamics 365 on that page.")
para("How that publish is done, as documented for CatalogONE [A][P][T][S]:")
table([
    ["Mechanism", "What Amdocs / TM Forum say", "Named target"],
    ["CatalogONE Centralized",
     "Author once; publish to multiple consuming applications [A]",
     "Consuming applications (not named as Dynamics)"],
    ["Kafka distribution",
     "Synchronise catalog definitions to other systems in minutes [P]",
     "Federated / downstream systems"],
    ["TMF Open APIs (TMF620 and related)",
     "API-first, TMF data model; Amdocs is certified on TMF620 [A][T]",
     "Any system that consumes TMF620"],
    ["CPQ + CRM",
     "\"integration to Amdocs CPQ and the customer's CRM solution\" [A]",
     "The operator's CRM (unnamed)"],
], widths=[4.4, 7.4, 4.8])

h2("4.2 What the Amdocs–Microsoft platform documents: Catalog next to Dynamics 365")
para("CEP is the Amdocs–Microsoft Customer Engagement Platform. It is not CatalogONE. "
     "It is a packaged telecom CRM that Amdocs and Microsoft sell together: "
     "Amdocs catalog, CPQ, commerce and order management, plus Microsoft Dynamics 365 "
     "Sales, Marketing and Customer Service, on Azure [A]. \"Buy CEP\" means buy that "
     "whole platform — not a CatalogONE plug-in for an existing Dynamics org.")
para("Amdocs' wording: the platform \"leverages … a unified Catalog, CPQ, Commerce and "
     "Order Management, and adds Microsoft's Dynamics 365 modules focused on sales, "
     "marketing, and customer service\" [A]. AppSource lists it as a Dynamics 365 "
     "offering, with \"Unified Catalog Serving the Entire Platform\" [A].")
para("In CEP papers Amdocs says \"Catalog\", not \"CatalogONE\". CatalogONE is "
     "Amdocs' catalog/PLM product, so it is the natural catalog in that stack, "
     "but that identification is not written as \"CatalogONE → Dynamics 365\" [A][E].")
image("01_cep_architecture.png",
      "Figure 4. CEP architecture: Catalog is its own module, next to Sales. "
      "In this platform Sales is Dynamics 365. Source: Amdocs and Microsoft "
      "Customer Engagement Platform paper, p.3 [A].", width_cm=16)
image("02_cep_catalog_and_crm.png",
      "Figure 5. Complementary stack: Digital Enterprise Catalog (Amdocs) and "
      "Sales management (Dynamics) on one platform. Source: Amdocs–Microsoft "
      "CEP paper, p.4 [A].", width_cm=16)
para("CEP also states that agent actions are \"captured and updated across all "
     "relevant billing, CRM, operations, and product databases\" [A]. That is "
     "platform-level sync, not a CatalogONE API named in a CatalogONE datasheet.")
para("A live operator on CEP is PLDT (Philippines). Amdocs announced the go-live "
     "on 7 May 2025. Agents use Microsoft Copilot and Microsoft Dynamics 365, "
     "with network validation and real-time dashboards [A]. In June 2026 IDC "
     "published a case study (Christopher Silberberg and Ahmad Latif Ali) on "
     "PLDT Unified Case Management on CEP: six weeks after go-live, 48% faster "
     "L1 ticket resolution and a 27-point NPS increase on the targeted use "
     "cases [A][I].")
para("That IDC paper is customer care: ticketing and case handling on CEP, with "
     "Dynamics 365 as the agent workspace. The May 2025 go-live also mentions "
     "faster service rollouts at platform level. Neither source describes "
     "CatalogONE publishing product-catalog data into Dynamics 365 Sales [A][I].")

h2("4.3 What is not publicly documented")
bullets([
    "A CatalogONE product page, datasheet or API guide that names Microsoft "
    "Dynamics 365 as the destination.",
    "An AppSource app called CatalogONE for Dynamics.",
    "A vendor field mapping from CatalogONE / TMF620 entities onto Dataverse "
    "Product, Price list item and Bundle.",
    "Whether CatalogONE can be plugged into an existing Dynamics 365 Sales org "
    "(such as WE Data) without buying the full Customer Engagement Platform.",
    "Any sentence in the IDC PLDT case study that says CatalogONE sends "
    "catalog data to Dynamics 365. That paper is Unified Case Management on CEP.",
])

h2("4.4 Example only: how TMF620 items would land in Dynamics 365 Sales [E][M]")
para("If CatalogONE (or the CEP catalog) is the source of truth, and Dynamics 365 "
     "Sales is the CRM, the industry mapping is the same pattern used for any "
     "TMF620 catalog. The table is an illustration from Microsoft's product-catalog "
     "entities plus the TMF model. It is not copied from an Amdocs CatalogONE "
     "connector guide [E][M][S].")
table([
    ["Catalog / TMF620 item", "Dynamics 365 Sales entity [M]"],
    ["Product Offering", "Product"],
    ["Bundle offering", "Product bundle"],
    ["Product Catalog / Category", "Product family"],
    ["Product Offering Price", "Price list item (within a price list)"],
    ["Discounts / promotions", "Discount list"],
    ["Related offers / add-ons", "Product relationships (accessory, cross-sell, up-sell)"],
    ["Characteristics", "Product properties (DynamicProperty)"],
    ["Lifecycle status: Active / Launched / Retired", "Product status: Active / Retired"],
], widths=[7.6, 9.0])
para("Dynamics product statuses are Draft, Active, Under Revision and Retired [M]. "
     "Study, design, test and approval would stay in CatalogONE; Dynamics would "
     "receive only products that are ready to sell [S][M][E].", italic=True)

h2("4.5 Two practical paths, if Dynamics is already the CRM")
table([
    ["Path", "What it is", "Evidence"],
    ["Buy the Customer Engagement Platform (CEP)",
     "The Amdocs–Microsoft packaged CRM: Amdocs catalog + Dynamics 365 as one "
     "platform. PLDT is live on CEP customer service (Unified Case Management)",
     "Amdocs–Microsoft papers [A]; PLDT go-live and IDC case study [A][I]"],
    ["Keep current Dynamics, add CatalogONE",
     "CatalogONE publishes via TMF620 / Kafka; an integration layer maps into "
     "Dataverse. This is how CatalogONE describes any CRM, not a boxed Dynamics connector",
     "CatalogONE: consuming applications + customer's CRM [A]; mapping is [E]"],
], widths=[4.2, 8.0, 4.4])
para("CatalogONE is documented as requiring \"no system or catalog replacements\" "
     "and as able to sit on top of existing catalogs and BSS [A]. That supports "
     "keeping Dynamics. It does not, by itself, prove a ready-made CatalogONE-to-"
     "Dynamics feed.")

# ------------------------------------------------------------------ 5
h1("5. Evidence from real operators")
table([
    ["Operator", "What was deployed", "Published result", "Source type"],
    ["European quad-play operator",
     "CatalogONE, unifying siloed lines of business",
     "New-offer time from up to 9 months to 1–2 weeks; changes to existing offers "
     "in hours; 80% of offer development by configuration (was 70% system integration)",
     "Appledore Research solution profile, 2020 [P]"],
    ["Sprint / T-Mobile US",
     "CatalogONE (define, implement, test, launch, manage offers) + DigitalONE",
     "Expected 50% time-to-market reduction; business users define offers "
     "(3–4 people instead of 20–30); 22 channel catalogs to one",
     "TM Forum Excellence Awards case study [T]"],
    ["KT Corporation (Korea)",
     "CatalogONE on KT private cloud, migrated charging product catalog",
     "Business-driven UI so teams configure changes directly; foundation to "
     "accelerate 5G consumer and enterprise offerings",
     "Amdocs press release, Nov 2020 [A]"],
    ["Mtel (A1 Bulgaria)",
     "Amdocs Master Enterprise Catalog (CatalogONE predecessor)",
     "Time to market for new offers cut by more than 80%; complex offers from "
     "idea to production in under three days (CTO quote)",
     "Operator quote in trade press, 2015"],
    ["PLDT (Philippines)",
     "Amdocs–Microsoft Customer Engagement Platform: Unified Case Management. "
     "Agents use Dynamics 365. Not a CatalogONE catalog-publish case",
     "Six weeks after go-live: 48% faster L1 resolution; +27 NPS on targeted "
     "use cases. Press names Dynamics 365 and Copilot as the agent workspace",
     "Amdocs go-live, 7 May 2025 [A]; IDC case study, June 2026 [I]"],
], widths=[3.4, 4.4, 5.2, 3.6])
image("10_sprint_results.png",
      "Figure 6. Sprint / T-Mobile US: CatalogONE as the product catalog for a "
      "cloud-native commerce transformation. Expected 50% time-to-market reduction; "
      "business users define offers. Source: TM Forum Excellence Awards case study [T].",
      width_cm=12)

# ------------------------------------------------------------------ 6
h1("6. Points to explore further")
bullets([
    "How ideas and business cases are captured and scored before In Study.",
    "How the approval workflow is configured: approvers per stage, sequential or "
    "parallel, delegation.",
    "Which roles and permissions exist for catalog users (view, create, edit, "
    "approve, publish).",
    "Whether custom lifecycle statuses can be added (for example a regulatory "
    "approval step).",
    "How versions are compared, and whether two versions can be in design at the "
    "same time.",
    "The exact CEP connector mapping from CatalogONE entities onto Dataverse "
    "product, price list and bundle tables.",
    "How failed synchronisations to Dynamics are retried and audited.",
    "How contract-term and multi-site pricing for enterprise offers are modelled "
    "when the Charging catalog is not in use.",
])

# ------------------------------------------------------------------ 7
h1("7. References")
refs = [
    "Amdocs, CatalogONE product page: "
    "https://www.amdocs.com/products-services/bss-oss/catalog-one",
    "Amdocs, CatalogONE – a new catalog paradigm (datasheet, 2025): "
    "https://solutions.amdocs.com/rs/647-OJR-802/images/Amdocs-catalog-datasheet-paradigm-02-2025.pdf",
    "Amdocs, CatalogONE – your business growth is as easy as the product it can launch (datasheet, 2025): "
    "https://solutions.amdocs.com/rs/647-OJR-802/images/Amdocs-catalog-datasheet-02-2025.pdf",
    "Amdocs, The 6 Steps to Turning Your Catalog into a Driver of Growth (2025): "
    "https://solutions.amdocs.com/rs/647-OJR-802/images/Amdocs-signaturecx-campaign-ebook-catalog-02-2025.pdf",
    "Amdocs, CatalogONE whitepaper (role-based UI, TMF APIs, insertion points): "
    "http://www.amdocs-marketing.com/CatalogONEwhitepaper/page-1.html",
    "Appledore Research, Amdocs CatalogONE Solution Profile (April 2020): "
    "https://www.amdocs.com/sites/default/files/2021-07/Amdocs-CatalogONE-Appledore-Solution-Profile.pdf",
    "Amdocs and Microsoft, Customer Engagement Platform paper (2025): "
    "https://solutions.amdocs.com/rs/647-OJR-802/images/Amdocs-and-microsoft-customer-engagement-platform-01-2025.pdf",
    "Amdocs and Microsoft, Customer Engagement Platform datasheet (2024): "
    "https://solutions.amdocs.com/rs/647-OJR-802/images/customer-engagement-platform-datasheet-04-11-2024.pdf",
    "Microsoft AppSource, Customer Engagement Platform (Amdocs): "
    "https://appsource.microsoft.com/en-us/product/dynamics-365/cyprus1676381115144.customerengagementplatform",
    "Amdocs, press release, Amdocs and Microsoft introduce Customer Engagement Platform (27 Feb 2023): "
    "https://www.amdocs.com/news-press/amdocs-and-microsoft-expand-strategic-partnership-reimagine-telco-experience",
    "Amdocs, press release, PLDT Goes Live with the Amdocs Customer Engagement Platform (7 May 2025): "
    "https://www.amdocs.com/press-release/pldt-goes-live-amdocs-customer-engagement-platform-redefine-customer",
    "IDC (Christopher Silberberg and Ahmad Latif Ali), How PLDT Delivered Customer "
    "Experience Gains in Just Six Weeks — Unified Case Management on the Amdocs "
    "Customer Engagement Platform (June 2026): "
    "https://www.amdocs.com/sites/default/files/2026-06/Amdocs-PLDT-idc-case-study.pdf",
    "Amdocs, case-study landing page for the same IDC paper: "
    "https://www.amdocs.com/insights/case-study/how-pldt-delivered-customer-experience-gains-just-six-weeks",
    "TM Forum, Sprint / T-Mobile cloud-native case study (CatalogONE): "
    "https://www.tmforum.org/wp-content/uploads/2021/02/TM-Forum-Excellence-Awards-Case-Study-Cloud_Sprint.pdf",
    "Amdocs, KT implements CatalogONE (10 Nov 2020): "
    "https://www.globenewswire.com/news-release/2020/11/10/2124219/0/en/"
    "KT-Successfully-Implements-Amdocs-CatalogONE-Cloud-Native-Solution-to-Accelerate-Introduction-of-New-5G-Offerings.html",
    "TM Forum, Open API conformance list (Amdocs TMF620 and related APIs): "
    "https://www.tmforum.org/conformance-certification/open-api-conformance/",
    "TM Forum, TMF620 Product Catalog Management API: "
    "https://www.tmforum.org/open-digital-architecture/open-apis/product-catalog-management-api-TMF620/v5.0",
    "TM Forum, TMF679 Product Offering Qualification API: "
    "https://www.tmforum.org/open-digital-architecture/open-apis/product-offering-qualification-api-TMF679",
    "Microsoft Learn, Product catalog tables (Dynamics 365 Sales): "
    "https://learn.microsoft.com/en-us/dynamics365/sales/developer/product-catalog-entities",
    "Microsoft Learn, Publish a product or bundle (Dynamics 365 Sales): "
    "https://learn.microsoft.com/en-us/dynamics365/sales/publish-product-bundle-make-available-selling",
]
numbered(refs)

doc.save(OUT)
print("saved", OUT.name)
